# 1-mavzu (M1, d001): ustoz bilan sinash uchun taqdimot (qoralama).
# Ishlatish: python3 tool/slides/make_m1.py chiqish.pptx  (python-pptx, Pillow)
# Mazmun: toifa og'zaki rejalari kdl-o-006/007/037, JSST qo'llanmalari (LQMS,
# Biosafety Manual, Hand Hygiene), CDC SHHV tartibi; tasdiqlanmagan — ustoz
# tekshiradi. Belgilar: Material Icons (Apache 2.0, Flutter SDK shrifti).
import json
import os
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
LOGO = os.path.join(REPO, 'assets', 'images', 'logo_mark.png')
ICON_FONT = os.path.join(os.environ.get('FLUTTER_ROOT', '/root/sdk/flutter'),
                         'bin/cache/artifacts/material_fonts/MaterialIcons-Regular.otf')
ICON_DIR = os.path.join(os.environ.get('TMPDIR', '/tmp'), 'lg_slide_icons')
FONT = 'Inter'

INK = RGBColor(0x18, 0x3A, 0x36)
DEEP = RGBColor(0x12, 0x2D, 0x2A)
BG = RGBColor(0xF4, 0xF2, 0xEC)
SOFT = RGBColor(0xE6, 0xEC, 0xDD)
LEAF = RGBColor(0x2E, 0x6B, 0x5E)
MUTED = RGBColor(0x53, 0x6C, 0x65)
ACC = RGBColor(0xC2, 0x7A, 0x2C)
RED = RGBColor(0xB4, 0x3A, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xC9, 0xD8, 0xC2)

# Material Icons kod nuqtalari (Flutter Icons.*).
I = dict(
    biotech=0xe0df, science=0xe55a, fire=0xe392, extinguisher=0xe288,
    electric=0xe228, masks=0xe3d3, safety=0xe305, warning=0xe6cb,
    verified=0xe69a, lock=0xe3ae, phone_off=0xe4a6, groups=0xe2ee,
    school=0xe559, doc=0xe1bf, inventory=0xe34a, sanitizer=0xe54e,
    virus=0xe199, blood=0xe0e3, vaccine=0xf0597, medical=0xe3d8,
    emergency=0xf04fa, hands=0xe166, hand=0xf0510, qr=0xe4f6,
    truck=0xe3a6, thermo=0xe656, check=0xe256, problem=0xe52d,
    building=0xe089, handshake=0xf06be, mind=0xe4ef, heart=0xe25b,
    gavel=0xe2d3, book=0xe3dd, voice=0xe510, list=0xe15b, delete=0xe1b9,
    timer=0xe662, person_off=0xe496, no_touch=0xe1f0, no_food=0xe438,
    no_smoke=0xe5c7, eye_off=0xe6be, eye=0xe6bd, coat=0xe15d, glove=0xe46c,
)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]
DATA = []  # ilova uchun (assets/lectures/d001.json)


def icon(name, color=INK, slash=False, px=256):
    """Belgini shaffof PNG qilib chizadi (keshlanadi)."""
    os.makedirs(ICON_DIR, exist_ok=True)
    rgb = (color[0], color[1], color[2])
    path = os.path.join(ICON_DIR, f'{name}_{color}_{int(slash)}.png')
    if os.path.exists(path):
        return path
    font = ImageFont.truetype(ICON_FONT, int(px * 0.86))
    im = Image.new('RGBA', (px, px), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    ch = chr(I[name])
    l, t, r, b = d.textbbox((0, 0), ch, font=font)
    d.text(((px - (r - l)) / 2 - l, (px - (b - t)) / 2 - t), ch, font=font, fill=rgb + (255,))
    if slash:
        w = int(px * 0.07)
        d.ellipse([w, w, px - w, px - w], outline=(180, 58, 46, 255), width=w)
        d.line([px * 0.2, px * 0.2, px * 0.8, px * 0.8], fill=(180, 58, 46, 255), width=w)
    im.save(path)
    return path


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def box(slide, x, y, w, h, color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = fill(slide.shapes.add_shape(shape, int(x), int(y), int(w), int(h)), color)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def pic(slide, name, x, y, size, color=INK, slash=False):
    slide.shapes.add_picture(icon(name, color, slash), int(x), int(y), int(size), int(size))


def badge(slide, name, x, y, d, bg=SOFT, fg=INK, slash=False):
    """Doira ichida belgi."""
    box(slide, x, y, d, d, bg, MSO_SHAPE.OVAL)
    pad = d * 0.2
    pic(slide, name, x + pad, y + pad, d - 2 * pad, fg, slash)


def text(slide, x, y, w, h, s, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0):
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(s if isinstance(s, list) else [s]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT
    return tb


def base(slide, title=None):
    box(slide, 0, 0, W, H, BG, MSO_SHAPE.RECTANGLE)
    slide.shapes.add_picture(LOGO, Inches(0.6), Inches(0.42), height=Inches(0.42))
    text(slide, Inches(1.12), Inches(0.46), Inches(6), Inches(0.4), 'LabGuide · 1-mavzu', 13, MUTED, True)
    text(slide, Inches(0.6), Inches(6.95), Inches(9), Inches(0.35),
         'Ustoz tekshiruvi uchun qoralama · LabGuide', 11, MUTED)
    text(slide, Inches(10.2), Inches(6.95), Inches(2.5), Inches(0.35),
         f'{len(prs.slides)}', 12, MUTED, align=PP_ALIGN.RIGHT)
    if title:
        text(slide, Inches(0.6), Inches(1.0), Inches(12.1), Inches(0.9), title, 30, INK, True)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def record(title, bullets=(), note='', source=None, eyebrow=None, section=False):
    d = dict(title=title, bullets=list(bullets))
    if eyebrow:
        d['eyebrow'] = eyebrow
    if note:
        d['note'] = note
    if source:
        d['source'] = source
    if section:
        d['section'] = True
    DATA.append(d)


def footnote(slide, foot):
    if foot:
        text(slide, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.4), foot, 12, MUTED)


# ---------------------------------------------------------------- shablonlar
def title_slide():
    s = prs.slides.add_slide(blank)
    box(s, 0, 0, W, H, INK, MSO_SHAPE.RECTANGLE)
    box(s, 0, Inches(6.6), W, Inches(0.9), DEEP, MSO_SHAPE.RECTANGLE)
    # O'ng tomonda belgi kompozitsiyasi.
    box(s, Inches(8.7), Inches(1.2), Inches(4.2), Inches(4.2), LEAF, MSO_SHAPE.OVAL)
    pic(s, 'biotech', Inches(9.55), Inches(2.05), Inches(2.5), WHITE)
    badge(s, 'safety', Inches(8.35), Inches(4.3), Inches(1.3), ACC, WHITE)
    badge(s, 'handshake', Inches(11.95), Inches(1.0), Inches(1.15), PALE, INK)
    s.shapes.add_picture(LOGO, Inches(0.8), Inches(0.8), height=Inches(0.8))
    text(s, Inches(1.8), Inches(0.92), Inches(6), Inches(0.6), 'LabGuide', 26, WHITE, True)
    text(s, Inches(0.8), Inches(2.0), Inches(7.6), Inches(0.5),
         'KLINIK LABORATOR DIAGNOSTIKA · 1-MODUL · 1-MAVZU', 14, PALE, True)
    text(s, Inches(0.8), Inches(2.55), Inches(7.6), Inches(2.4),
         ['Laborator xizmatni', 'tashkillashtirish.', 'Tibbiy etika va deontologiya'], 38, WHITE, True)
    text(s, Inches(0.8), Inches(5.05), Inches(7.6), Inches(0.9),
         'Amaliy mashg‘ulot (4 soat) + Ma’naviyat darsi (2 soat)', 18, RGBColor(0xDD, 0xE6, 0xD8))
    text(s, Inches(0.8), Inches(6.82), Inches(11.8), Inches(0.5),
         'Ustoz tekshiruvi uchun qoralama — dars oldidan mazmunni ko‘rib chiqing', 14, PALE)
    notes(s, 'Kirish. Dars rejasi: laboratoriya xizmati tuzilishi, namuna yo‘li, sifat tizimi, xavfsizlik '
             '(umumiy, yong‘in, biologik, qo‘l gigiyenasi, SHHV), etika va deontologiya, ma’naviyat. '
             'Oxirida mini-test va og‘zaki savollar (toifa banki: kdl-o-006, kdl-o-007, kdl-o-037).')
    record('Laborator xizmatni tashkillashtirish. Tibbiy etika va deontologiya',
           ['Amaliy mashg‘ulot (4 soat) + Ma’naviyat darsi (2 soat)'], eyebrow='1-modul · 1-mavzu')


def section(n, t, sub, ic):
    s = prs.slides.add_slide(blank)
    box(s, 0, 0, W, H, INK, MSO_SHAPE.RECTANGLE)
    box(s, Inches(8.9), Inches(1.6), Inches(3.8), Inches(3.8), LEAF, MSO_SHAPE.OVAL)
    pic(s, ic, Inches(9.75), Inches(2.45), Inches(2.1), WHITE)
    text(s, Inches(0.8), Inches(2.0), Inches(3), Inches(1.4), f'{n:02d}', 80, ACC, True)
    text(s, Inches(0.8), Inches(3.5), Inches(7.8), Inches(1.6), t, 36, WHITE, True)
    text(s, Inches(0.8), Inches(5.0), Inches(7.8), Inches(0.8), sub, 19, PALE)
    record(t, [sub], eyebrow=f'{n}-bo‘lim', section=True)


def bullets(t, items, ic, note='', foot=None, caption=None, slash=False):
    """Chapda bandlar, o'ngda illyustratsiya paneli."""
    s = prs.slides.add_slide(blank)
    base(s, t)
    box(s, Inches(9.35), Inches(2.0), Inches(3.4), Inches(4.3), SOFT)
    badge(s, ic, Inches(9.85), Inches(2.45), Inches(2.4), WHITE, INK, slash)
    if caption:
        text(s, Inches(9.55), Inches(5.1), Inches(3.0), Inches(1.0), caption, 14, MUTED,
             True, PP_ALIGN.CENTER)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(8.4), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        head, _, rest = it.partition(' — ')
        for txt, sz, b, c in [('●  ', 13, False, ACC), (head, 19, bool(rest), INK),
                              ((' — ' + rest) if rest else '', 19, False, INK)]:
            if not txt:
                continue
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(sz)
            r.font.bold = b
            r.font.color.rgb = c
            r.font.name = FONT
    footnote(s, foot)
    notes(s, note)
    record(t, items, note, foot)
    return s


def cards(t, items, cols=3, note='', foot=None):
    """items: (belgi, sarlavha, matn[, slash])."""
    s = prs.slides.add_slide(blank)
    base(s, t)
    rows = (len(items) + cols - 1) // cols
    gap = Inches(0.25)
    cw = (Inches(12.13) - gap * (cols - 1)) / cols
    top, avail = Inches(2.0), Inches(4.35)
    ch = (avail - gap * (rows - 1)) / rows
    small = cols > 3 or rows > 2
    d = Inches(0.62) if small else Inches(0.85)
    for i, it in enumerate(items):
        ic, h, b = it[:3]
        slash = len(it) > 3 and it[3]
        r, c = divmod(i, cols)
        x = Inches(0.6) + (cw + gap) * c
        y = top + (ch + gap) * r
        box(s, x, y, cw, ch, SOFT)
        badge(s, ic, x + Inches(0.22), y + Inches(0.2), d, WHITE, INK, slash)
        tx = x + Inches(0.22) + d + Inches(0.18)
        text(s, tx, y + Inches(0.22), cw - (tx - x) - Inches(0.2), d, h,
             15 if small else 19, INK, True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + Inches(0.22), y + Inches(0.32) + d, cw - Inches(0.44),
             ch - d - Inches(0.4), b, 13 if small else 16, MUTED)
    footnote(s, foot)
    notes(s, note)
    record(t, [f'{it[1]} — {it[2]}'.replace('\n', '; ') for it in items], note, foot)
    return s


def flow(t, steps, note='', foot=None, top=Inches(2.15), height=Inches(3.9), hi=1):
    """steps: (belgi, sarlavha, matn). Ajratilgan qadam — [hi]."""
    s = prs.slides.add_slide(blank)
    base(s, t)
    n = len(steps)
    gap = Inches(0.4)
    cw = (Inches(12.13) - gap * (n - 1)) / n
    for i, (ic, h, b) in enumerate(steps):
        x = Inches(0.6) + (cw + gap) * i
        dark = i == hi
        box(s, x, top, cw, height, INK if dark else SOFT, radius=0.06)
        badge(s, ic, x + Inches(0.3), top + Inches(0.3), Inches(0.9),
              LEAF if dark else WHITE, WHITE if dark else INK)
        text(s, x + Inches(1.35), top + Inches(0.42), cw - Inches(1.5), Inches(0.4),
             f'{i + 1}-qadam', 13, ACC, True)
        text(s, x + Inches(0.3), top + Inches(1.35), cw - Inches(0.6), Inches(0.8),
             h, 21, WHITE if dark else INK, True)
        text(s, x + Inches(0.3), top + Inches(2.05), cw - Inches(0.6), height - Inches(2.2),
             b, 15, PALE if dark else MUTED)
        if i < n - 1:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, int(x + cw + Inches(0.06)),
                                   int(top + height / 2 - Inches(0.18)), Inches(0.28), Inches(0.36))
            fill(a, ACC)
    footnote(s, foot)
    notes(s, note)
    record(t, [f'{i + 1}. {h} — {b}' for i, (_, h, b) in enumerate(steps)], note, foot)
    return s


def chain(t, steps, note='', foot=None):
    """Belgili doiralar zanjiri (namuna yo'li)."""
    s = prs.slides.add_slide(blank)
    base(s, t)
    n = len(steps)
    d = Inches(1.05)
    span = Inches(12.13)
    step = (span - d) / (n - 1)
    cy = Inches(2.75)
    line = s.shapes.add_connector(1, int(Inches(0.6) + d / 2), int(cy + d / 2),
                                  int(Inches(0.6) + span - d / 2), int(cy + d / 2))
    line.line.color.rgb = PALE
    line.line.width = Pt(4)
    for i, (ic, h, b) in enumerate(steps):
        x = Inches(0.6) + step * i
        badge(s, ic, x, cy, d, INK if i in (1, 3) else SOFT, WHITE if i in (1, 3) else INK)
        text(s, x - Inches(0.25), Inches(1.95), d + Inches(0.5), Inches(0.5), f'{i + 1}', 16, ACC,
             True, PP_ALIGN.CENTER)
        text(s, x - Inches(0.3), cy + d + Inches(0.2), d + Inches(0.6), Inches(0.7), h, 15, INK,
             True, PP_ALIGN.CENTER)
        text(s, x - Inches(0.3), cy + d + Inches(0.85), d + Inches(0.6), Inches(1.6), b, 12, MUTED,
             align=PP_ALIGN.CENTER)
    footnote(s, foot)
    notes(s, note)
    record(t, [f'{i + 1}. {h} — {b}' for i, (_, h, b) in enumerate(steps)], note, foot)
    return s


def ppe(t, don, doff, note='', foot=None):
    """SHHV: kiyish va yechish tartibi — ikki qator."""
    s = prs.slides.add_slide(blank)
    base(s, t)
    for row, (label, steps, color) in enumerate([('Kiyish', don, LEAF), ('Yechish', doff, ACC)]):
        y = Inches(2.0) + row * Inches(2.25)
        box(s, Inches(0.6), y, Inches(1.7), Inches(1.95), color)
        text(s, Inches(0.6), y, Inches(1.7), Inches(1.95), label, 20, WHITE, True,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        n = len(steps)
        gap = Inches(0.18)
        cw = (Inches(10.25) - gap * (n - 1)) / n
        for i, (ic, h) in enumerate(steps):
            x = Inches(2.48) + (cw + gap) * i
            box(s, x, y, cw, Inches(1.95), SOFT)
            text(s, x + Inches(0.12), y + Inches(0.1), Inches(0.5), Inches(0.4), f'{i + 1}', 15, color, True)
            badge(s, ic, x + cw / 2 - Inches(0.42), y + Inches(0.22), Inches(0.84), WHITE, INK)
            text(s, x + Inches(0.08), y + Inches(1.15), cw - Inches(0.16), Inches(0.75), h, 13, INK,
                 True, PP_ALIGN.CENTER)
    footnote(s, foot)
    notes(s, note)
    record(t, ['Kiyish: ' + ' → '.join(h for _, h in don), 'Yechish: ' + ' → '.join(h for _, h in doff)],
           note, foot)
    return s


def quiz(t, qs, note='', foot=None):
    """qs: (savol, [variantlar]). Javoblar faqat ustoz izohida."""
    s = prs.slides.add_slide(blank)
    base(s, t)
    gap = Inches(0.25)
    cw = (Inches(12.13) - gap) / 2
    ch = (Inches(4.35) - gap) / 2
    for i, (q, opts) in enumerate(qs):
        r, c = divmod(i, 2)
        x = Inches(0.6) + (cw + gap) * c
        y = Inches(2.0) + (ch + gap) * r
        box(s, x, y, cw, ch, SOFT)
        badge(s, 'list', x + Inches(0.2), y + Inches(0.2), Inches(0.55), WHITE, INK)
        text(s, x + Inches(0.9), y + Inches(0.18), cw - Inches(1.1), Inches(0.9), f'{i + 1}. {q}', 16, INK, True)
        text(s, x + Inches(0.9), y + Inches(1.0), cw - Inches(1.1), ch - Inches(1.1),
             [f'{"ABCD"[k]})  {o}' for k, o in enumerate(opts)], 14, MUTED, space=2)
    footnote(s, foot)
    notes(s, note)
    record(t, [f'{i + 1}. {q}' for i, (q, _) in enumerate(qs)], note, foot)
    return s


# ---------------------------------------------------------------- slaydlar
title_slide()

bullets('Dars maqsadi va rejasi', [
    'Laboratoriya xizmati — vazifalari, tuzilmasi, namuna yo‘li',
    'Sifat tizimi — JSST LQMS ning 12 ta elementi, ISO 15189',
    'Xavfsizlik — umumiy, yong‘in, biologik; qo‘l gigiyenasi va SHHV',
    'Tibbiy etika va deontologiya — tamoyillar, konfidensiallik, halollik',
    'Ma’naviyat darsi — kasbiy burch, jamoa, bemor oldidagi mas’uliyat',
    'Yakun — vaziyatlar, mini-test va og‘zaki savollar',
], 'list', 'Talabalar dars oxirida: tahlil jarayoni bosqichlarini, sifat tizimi elementlarini, '
           'asosiy xavfsizlik qoidalarini va etika tamoyillarini aytib bera olishi kerak.',
   caption='6 soatlik dars')

section(1, 'Laborator xizmatni tashkillashtirish', 'Laboratoriya — tashxis jarayonining ajralmas qismi', 'building')

cards('Klinik laboratoriyaning vazifalari', [
    ('science', 'Tashxis', 'Kasallikni aniqlash va tasdiqlash uchun ishonchli natija'),
    ('check', 'Monitoring', 'Davolash samarasini va kasallik kechishini kuzatish'),
    ('safety', 'Skrining', 'Yashirin kasalliklarni erta aniqlash, xavf omillarini baholash'),
    ('emergency', 'Shoshilinch yordam', 'Kritik natijalarni o‘z vaqtida shifokorga yetkazish'),
    ('verified', 'Sifat va xavfsizlik', 'Natijalar aniqligi, xodim va bemor xavfsizligi'),
    ('school', 'Ta’lim va ilm', 'Xodimlarni o‘qitish, yangi metodlarni joriy etish'),
], 3, 'Laboratoriya natijasi ko‘plab klinik qarorlarga asos bo‘ladi — shuning uchun xato narxi yuqori. '
      'Misol: kaliy natijasi noto‘g‘ri bo‘lsa (gemoliz), noto‘g‘ri davolash boshlanishi mumkin.')

cards('Laboratoriya jamoasi: kim nima qiladi', [
    ('building', 'Laboratoriya mudiri', 'Ishni tashkil etish, sifat tizimi va xavfsizlik uchun javobgar'),
    ('medical', 'KDL shifokori', 'Natijani klinik baholash va tasdiqlash, shifokorlar bilan maslahat'),
    ('biotech', 'Laborant', 'Namunani qabul qilish, tahlil bajarish, QC, uskunaga qarash'),
    ('sanitizer', 'Kichik tibbiy xodim', 'Tozalik, dezinfeksiya, chiqindilarni tartib bilan yig‘ish'),
], 2, 'Umumiy rollar. Aniq shtat birliklari va lavozim yo‘riqnomalari — muassasa va amaldagi me’yoriy hujjatlar '
      'bo‘yicha: ustoz o‘z muassasasidagi tuzilmani misol qilib aytsin.')

flow('Tahlil jarayoni: uch bosqich', [
    ('blood', 'Preanalitik', 'Buyurtma, bemorni tayyorlash, namuna olish, belgilash, tashish, qabul'),
    ('science', 'Analitik', 'O‘lchash, kalibrlash, ichki sifat nazorati (QC), uskunaga xizmat'),
    ('doc', 'Postanalitik', 'Natijani tekshirish va tasdiqlash, hisobot, kritik qiymatni xabar qilish'),
], 'Adabiyotda laboratoriya xatolarining eng katta qismi preanalitik bosqichga to‘g‘ri kelishi '
   'ko‘rsatilgan (taxminan 60–70 %; Plebani M., Clin Chem Lab Med 2006). '
   'Ilovada: Lab → Preanalitika bo‘limi (namuna olish tartibi, JSST).',
   'Xatolarning katta qismi preanalitik bosqichda yuzaga keladi (Plebani, 2006).', hi=0)

chain('Namuna yo‘li: buyurtmadan natijagacha', [
    ('doc', 'Buyurtma', 'Shifokor yo‘llanmasi, tahlil turi'),
    ('hand', 'Identifikatsiya', 'Kamida 2 ta belgi: F.I.Sh., tug‘ilgan sana yoki ID'),
    ('blood', 'Namuna olish', 'To‘g‘ri probirka va tartib'),
    ('qr', 'Belgilash', 'Bemor oldida, darhol'),
    ('truck', 'Tashish', 'Vaqt va harorat sharti'),
    ('inventory', 'Qabul', 'Tekshirish, ro‘yxatga olish'),
    ('biotech', 'Tahlil', 'QC bilan'),
    ('check', 'Hisobot', 'Tasdiqlash, kritik qiymat'),
], 'Har bir bosqichda xato bo‘lishi mumkin. Talabalardan har bosqich uchun bitta mumkin bo‘lgan xatoni '
   'aytishni so‘rang (masalan: noto‘g‘ri bemor, noto‘g‘ri probirka, uzoq tashish, belgisiz namuna).',
   'Eng xavfli xato — noto‘g‘ri bemor: identifikatsiya har doim kamida 2 belgi bilan.')

bullets('Namunani belgilash va qabul qilish', [
    'Belgilash — kamida 2 ta identifikator, namuna olingan sana va vaqt, olgan xodim',
    'Qachon — probirka bemor oldida, namuna olingandan so‘ng darhol belgilanadi',
    'Rad etish — belgisiz yoki noto‘g‘ri belgilangan namuna',
    'Rad etish — noto‘g‘ri probirka, hajm yetarli emas, ivigan (antikoagulyantli probirkada)',
    'Rad etish — tashish sharti buzilgan (vaqt, harorat); gemoliz — tahlilga qarab',
    'Qayd etish — har bir rad etish hodisa sifatida yoziladi, bo‘limga xabar beriladi',
], 'qr', 'Rad etish mezonlari laboratoriyaning yozma tartibida (SOP) bo‘lishi kerak. '
         'Muhokama: “Belgisiz probirka keldi, lekin hamshira kimniki ekanini biladi” — qabul qilinadimi? (Yo‘q.)',
   caption='Ikki identifikator')

cards('Sifat tizimi — JSST LQMS ning 12 ta elementi', [
    ('building', 'Tashkilot', 'Tuzilma, rahbariyat'),
    ('groups', 'Xodimlar', 'Malaka, o‘qitish'),
    ('biotech', 'Uskunalar', 'Kalibrlash, xizmat'),
    ('inventory', 'Xarid va zaxira', 'Reagent, muddat'),
    ('science', 'Jarayon nazorati', 'QC, namuna'),
    ('lock', 'Axborot', 'Natija, maxfiylik'),
    ('doc', 'Hujjatlar', 'SOP, jurnallar'),
    ('problem', 'Hodisalar', 'Xatoni qayd etish'),
    ('check', 'Baholash', 'Audit, EQA'),
    ('verified', 'Yaxshilash', 'Tuzatuvchi choralar'),
    ('voice', 'Mijozga xizmat', 'Shifokor, bemor fikri'),
    ('safety', 'Bino, xavfsizlik', 'Joylashuv, himoya'),
], 4, 'Manba: WHO. Laboratory Quality Management System: Handbook. 2011. '
      'Xalqaro standart: ISO 15189:2022 — tibbiy laboratoriyalar sifati va kompetensiyasi.',
   'Manba: WHO LQMS Handbook (2011); ISO 15189:2022')

bullets('Hujjatlar va yozuvlar', [
    'SOP — har bir tahlil va jarayon uchun yozma, tasdiqlangan, ish joyida mavjud',
    'Jurnallar — QC natijalari, uskuna xizmati, harorat, reagent lotlari',
    'Kuzatuvchanlik — “kim, qachon, qaysi uskuna va reagent bilan” savoliga javob',
    'Hodisalar — xato va nomuvofiqliklar yashirilmaydi, qayd etiladi va tahlil qilinadi',
    'Qoida — “Yozilmagan ish — bajarilmagan ish”',
], 'doc', 'Amaliy topshiriq: guruh bitta oddiy tahlil (masalan, glyukoza) uchun SOP tuzilmasini yozsin: '
          'maqsad, namuna, reagent, uskuna, QC, hisoblash, referens oraliq, xavfsizlik.',
   caption='SOP va jurnallar')

section(2, 'Laboratoriya xavfsizligi', 'Xodim, bemor va atrof-muhitni himoya qilish', 'safety')

cards('Umumiy xavfsizlik tizimi', [
    ('list', 'Tashkiliy', 'Yozma xavfsizlik rejasi va SOP, mas’ul shaxs, ish boshlashda va davriy o‘qitish, favqulodda vaziyat mashqlari'),
    ('electric', 'Elektr', 'Shikastlangan rozetka/uskuna belgilanib ishdan chiqariladi; nam qo‘l bilan ulash taqiqlanadi; ta’mirda tokdan ajratish'),
    ('fire', 'Yong‘in', 'Yonuvchi erituvchilar minimal miqdorda, maxsus shkafda; ochiq olov nazoratsiz qoldirilmaydi; yo‘laklar ochiq'),
], 3, 'Toifa og‘zaki savoli kdl-o-006 rejasi asosida. Har bir xodim avariya jihozlari joyini bilishi kerak: '
      'o‘t o‘chirgich, signal, chiqish yo‘llari, avariya dushi/ko‘z yuvgich, birinchi yordam qutisi.')

cards('Yong‘in chiqqanda: RACE va PASS', [
    ('groups', 'RACE — harakat', 'Rescue — odamlarni olib chiqish\nAlarm — signal\nConfine — eshikni yopib cheklash\nEvacuate/Extinguish — chiqish yoki o‘chirish'),
    ('extinguisher', 'PASS — o‘t o‘chirgich', 'Pull — chekani tortish\nAim — olov asosiga\nSqueeze — richagni bosish\nSweep — yonma-yon purkash'),
    ('fire', 'O‘t o‘chirgich sinfi', 'A — qattiq yonuvchilar\nB — yonuvchi suyuqlik/gaz\nC — tok ostidagi uskuna\nKiyim yonsa — “to‘xta, yot, dumala”'),
], 3, 'kdl-o-006. Amaliy mashq: talabalar laboratoriyadagi o‘t o‘chirgich, avariya dushi va chiqish yo‘lini ko‘rsatib bersin.')

bullets('Biologik xavfsizlik: asoslar', [
    'Asos — har bir ish uchun xavfni baholash: ma’lumot → baholash → nazorat → choralar → qayta ko‘rish',
    'BSL darajalari — qon, tana suyuqliklari va to‘qimalar bilan ish odatda BSL-2 sharoitida',
    'Har bir namuna — potensial yuqumli deb qaraladi (standart ehtiyot choralari)',
], 'virus', 'kdl-o-007; manba: WHO Laboratory Biosafety Manual, 4th ed. (2020). '
            'BSL-1 — sog‘lom kattalarda kasallik chaqirmaydigan agentlar; BSL-2 — HBV, OIV, Salmonella va h.k.; '
            'BSL-3/4 — havo orqali yuqadigan og‘ir agentlar.',
   'Manba: WHO Laboratory Biosafety Manual, 4-nashr (2020)', caption='Har namuna — yuqumli')

cards('Laboratoriyada taqiqlanadi', [
    ('no_food', 'Ovqat va ichimlik', 'Ish zonasida ovqatlanish, ichish, ovqat saqlash'),
    ('no_smoke', 'Chekish', 'Chekish va kosmetika surtish'),
    ('science', 'Og‘iz bilan pipetkalash', 'Faqat mexanik dozator', True),
    ('no_touch', 'Qo‘lqopda “toza” yuza', 'Qo‘lqop bilan eshik, telefon, klaviatura ushlash'),
    ('person_off', 'Xalatda tashqariga', 'Xalat va qo‘lqop laboratoriyadan tashqarida kiyilmaydi'),
    ('masks', 'Himoyasiz ish', 'Sachrash xavfida ko‘z va yuz himoyasisiz ishlash'),
], 3, 'kdl-o-007 rejasi; WHO Laboratory Biosafety Manual (2020) — asosiy mikrobiologik amaliyot. '
      'Talabalar laboratoriyadagi ogohlantiruvchi belgilarni topib, ma’nosini aytsin.')

bullets('Qo‘l gigiyenasi', [
    'Spirtli antiseptik — 20–30 soniya, qo‘l quriguncha',
    'Sovun va suv — 40–60 soniya; qo‘l ko‘rinadigan darajada ifloslanganda',
    'Qachon — ish boshlashda va laboratoriyadan chiqishda',
    'Qachon — qo‘lqop kiyishdan oldin va yechgandan keyin',
    'Qachon — biomaterial bilan ishlagandan keyin, ovqatlanishdan oldin',
    'Eslatma — qo‘lqop qo‘l gigiyenasini almashtirmaydi',
], 'hands', 'Manba: WHO Guidelines on Hand Hygiene in Health Care (2009). Amaliy mashq: talabalar antiseptik bilan '
            'qo‘l ishlov berish texnikasini ko‘rsatsin (kaft, barmoqlar orasi, bosh barmoq, tirnoqlar, bilak).',
   'Manba: WHO Guidelines on Hand Hygiene in Health Care (2009)', caption='20–30 s / 40–60 s')

ppe('Shaxsiy himoya vositalari: kiyish va yechish tartibi', [
    ('hands', 'Qo‘l gigiyenasi'), ('coat', 'Xalat'), ('masks', 'Niqob'),
    ('eye', 'Ko‘zoynak / yuz himoyasi'), ('glove', 'Qo‘lqop'),
], [
    ('glove', 'Qo‘lqop'), ('eye', 'Ko‘zoynak / yuz himoyasi'), ('coat', 'Xalat'),
    ('masks', 'Niqob'), ('hands', 'Qo‘l gigiyenasi'),
], 'Manba: CDC — Sequence for putting on / removing PPE. Yechishda eng ifloslangan qism (qo‘lqop) birinchi '
   'yechiladi; tashqi yuzaga tegmaslik.',
   'Manba: CDC, Sequence for putting on and removing PPE. Yechish — eng iflosidan boshlab.')

bullets('Biologik xavfsizlik: amaliyot', [
    'O‘tkir buyumlar — igna qopqog‘i qayta kiydirilmaydi; darhol teshilmaydigan konteynerga',
    'Aerozol — sentrifugalash va ochish yopiq stakanlarda yoki biologik xavfsizlik shkafida',
    'To‘kilish — tayyor tartib bo‘yicha: cheklash, dezinfeksiya, yig‘ish',
    'Ish yuzalari — ish oxirida va har to‘kilishdan keyin dezinfeksiya',
], 'warning', 'kdl-o-007. Dekontaminatsiya: yuqumli material yo‘q qilishdan oldin (afzali avtoklavda) zararsizlantiriladi.',
   'Manba: WHO Laboratory Biosafety Manual, 4-nashr (2020)', caption='Aerozol va to‘kilish')

flow('Igna sanchilsa yoki qon shilliq qavatga tushsa', [
    ('hands', 'Darhol', 'Teri — oqar suv va sovun bilan yuvish; ko‘z/shilliq qavat — ko‘p suv bilan chayish'),
    ('voice', 'Xabar berish', 'Rahbarga darhol xabar berish, hodisani qayd etish'),
    ('medical', 'Profilaktika', 'Muassasa tartibi bo‘yicha ekspozitsiyadan keyingi baholash va profilaktika — imkon qadar tezroq'),
], 'kdl-o-007: xodimlar salomatligi — kasbiy sog‘liq talablari (jumladan, emlash), ekspozitsiyadan keyingi profilaktika tartibi. '
   'Aniq dori va muddatlar muassasa protokoli va amaldagi milliy ko‘rsatmaga ko‘ra — ustoz mahalliy tartibni aytib o‘tsin.',
   'Har bir laboratoriyada ekspozitsiyadan keyingi tartib yozma bo‘lishi va xodimlarga ma’lum bo‘lishi kerak.')

cards('Tibbiy chiqindilar', [
    ('warning', 'O‘tkir buyumlar', 'Igna, skarifikator, shisha — teshilmaydigan konteyner; ¾ to‘lganda yopiladi'),
    ('virus', 'Yuqumli', 'Biohazard belgili qop/idish; yo‘q qilishdan oldin zararsizlantirish (avtoklav)'),
    ('science', 'Kimyoviy', 'Reagent qoldiqlari — alohida, yorliq va xavfsizlik varag‘i bo‘yicha'),
    ('doc', 'Milliy talab', 'Chiqindi sinflari, rangi va tartibi — amaldagi sanitariya qoidalari bo‘yicha'),
], 2, 'Manba: WHO, Safe management of wastes from health-care activities (2014). Chiqindi sinflari va '
      'ranglari mamlakat sanitariya qoidalari bilan belgilanadi — ustoz amaldagi milliy tartibni aytsin.',
   'Manba: WHO, Safe management of wastes from health-care activities (2014)')

section(3, 'Tibbiy etika va deontologiya', 'Kasbiy burch va xulq-atvor', 'handshake')

cards('Tushunchalar va asosiy tamoyillar', [
    ('book', 'Tibbiy etika', 'Bemor, hamkasblar va jamiyatga munosabatdagi axloqiy tamoyillar'),
    ('gavel', 'Deontologiya', 'Kasbiy burch va xulq-atvor qoidalari haqidagi ta’limot'),
    ('heart', 'Bemor manfaati', 'Bemor manfaatini ustun qo‘yish, zarar yetkazmaslik'),
    ('hand', 'Avtonomiya', 'Bemor huquqi va ixtiyoriy xabardor rozilik'),
    ('groups', 'Kamsitmaslik', 'Xizmat faqat tibbiy ehtiyojga qarab'),
    ('verified', 'Adolat', 'Har bir bemorga bir xil sifat va e’tibor'),
], 3, 'kdl-o-037 rejasi asosida.')

bullets('Laboratoriyada konfidensiallik', [
    'Tibbiy sir — natijalar va bemor ma’lumotlari faqat vakolatli shaxslarga',
    'Muddatsiz — majburiyat bemor vafotidan keyin ham saqlanadi',
    'Telefon — natija telefon orqali begonaga aytilmaydi',
    'Blankalar — ochiq qoldirilmaydi, ekranlar qulflanadi',
    'Elektron tizim — shaxsiy login, parol hech kimga berilmaydi',
    'Ijtimoiy tarmoq — natija yoki namuna surati joylanmaydi',
], 'lock', 'Muhokama: “Qo‘shnim tanishining tahlil natijasini so‘radi” — laborant qanday javob beradi?',
   caption='Tibbiy sir')

cards('Kasbiy halollik va javobgarlik', [
    ('school', 'Kompetensiya', 'Doimiy o‘qish va malaka oshirish; bilmagan metodni mustaqil bajarmaslik'),
    ('verified', 'Halollik', 'Xato yashirilmaydi — qayd etiladi va tuzatiladi. Natijani “to‘g‘rilab” yozish taqiqlanadi'),
    ('emergency', 'Kritik natija', 'Hayot uchun xavfli qiymat darhol shifokorga yetkaziladi va qayd etiladi'),
    ('handshake', 'Hamkasblar', 'Hurmat va aniq kommunikatsiya; bemor oldida hamkasb tanqid qilinmaydi'),
], 2, 'kdl-o-037. O‘zbekistonda tibbiyot xodimlarining odob-axloq kodeksi SSV ning 2022-yil 24-martdagi 88-son buyrug‘i bilan '
      'tasdiqlangan — ustoz matnini lex.uz da tekshirib, asosiy talablarini o‘qib bersin.',
   'Tibbiyot xodimlari odob-axloq kodeksi (SSV, 2022, 88-son buyruq) — matnini lex.uz da tekshiring')

section(4, 'Ma’naviyat darsi', 'Kasb — bu mas’uliyat', 'heart')

bullets('Laborant — ko‘rinmas, lekin hal qiluvchi', [
    'Har bir probirka ortida — inson: bola, ona, keksa bemor',
    'Aniqlik — g‘amxo‘rlik: to‘g‘ri belgilangan namuna hayotni saqlaydi',
    'Jamoa — laboratoriya, klinika va bemor o‘rtasidagi ishonch',
    'O‘qishdan to‘xtamaslik — tibbiyot har yili yangilanadi',
    'Halollik — “Bilmayman, aniqlab beraman” deyish kasbiy jasorat',
], 'mind', 'Ustoz o‘z tajribasidan bitta voqea aytib berishi tavsiya etiladi. Muhokama: “Nega men bu kasbni tanladim?”',
   caption='Kasbiy burch')

bullets('Amaliy vaziyatlar (muhokama)', [
    '1 — Namuna belgisiz keldi. Nima qilasiz?',
    '2 — Kaliy 7,2 mmol/L, namuna gemolizlangan. Natijani berasizmi?',
    '3 — Hamkasbingiz QC “chiqib ketgan”da ham natija yubormoqchi.',
    '4 — Qo‘lingizga igna sanchildi. Birinchi 3 ta harakatingiz?',
    '5 — Bemorning qarindoshi natijani telefon orqali so‘radi.',
], 'voice', 'Kutilgan javoblar: 1) qabul qilinmaydi, qayta olish so‘raladi, hodisa qayd etiladi; '
            '2) gemoliz kaliyni soxta oshiradi — natija izoh bilan yoki namuna qayta olinadi, shifokorga xabar; '
            '3) QC talablariga javob bermaguncha bemor natijasi chiqarilmaydi — sabab topiladi, tuzatiladi; '
            '4) yuvish, rahbarga xabar, ekspozitsiyadan keyingi tartib; 5) vakolatsiz shaxsga natija berilmaydi.',
   caption='Guruhda muhokama')

quiz('Mini-test', [
    ('Laboratoriya xatolarining eng ko‘p qismi qaysi bosqichda?', ['Preanalitik', 'Analitik', 'Postanalitik']),
    ('Qon bilan ish odatda qaysi biologik xavfsizlik darajasida?', ['BSL-1', 'BSL-2', 'BSL-4']),
    ('SHHVni yechishda birinchi nima yechiladi?', ['Niqob', 'Xalat', 'Qo‘lqop']),
    ('Spirtli antiseptik bilan qo‘l ishlovi qancha davom etadi?', ['5 soniya', '20–30 soniya', '5 daqiqa']),
], 'Javoblar: 1 — A (preanalitik); 2 — B (BSL-2); 3 — C (qo‘lqop); 4 — B (20–30 soniya). '
   'Ilovada guruh testi: O‘rganish → Guruhlar → mavzu → test.')

bullets('Nazorat savollari (toifa — og‘zaki)', [
    'kdl-o-006 — Umumiy xavfsizlik, elektr va yong‘in xavfsizligi',
    'kdl-o-007 — Biologik xavfsizlik: xavfni baholash, BSL, SHHV, chiqindilar',
    'kdl-o-037 — Tibbiy etika va deontologiya, konfidensiallik',
    'Ilovada — O‘rganish → Toifa → Og‘zaki savollar: javob rejasi bilan',
], 'school', 'Savollar LabGuide toifa bankidan (rasmiy savollar). Javob rejalari ilovada mavjud.',
   caption='Toifa savollari')

s = prs.slides.add_slide(blank)
base(s, 'Manbalar')
SOURCES = [
    'WHO. Laboratory Quality Management System: Handbook. Geneva: WHO; 2011.',
    'WHO. Laboratory Biosafety Manual, 4th edition. Geneva: WHO; 2020.',
    'WHO. Guidelines on Hand Hygiene in Health Care. Geneva: WHO; 2009.',
    'WHO. Safe management of wastes from health-care activities, 2nd ed. Geneva: WHO; 2014.',
    'CDC. Sequence for putting on and removing personal protective equipment (PPE).',
    'ISO 15189:2022. Medical laboratories — Requirements for quality and competence.',
    'Plebani M. Errors in clinical laboratories or errors in laboratory medicine? Clin Chem Lab Med. 2006;44(6):750–759.',
    'O‘zR SSV, 2022-yil 24-mart, 88-son buyruq — tibbiyot xodimlari odob-axloq kodeksi (matni: lex.uz).',
    'LabGuide toifa banki: kdl-o-006, kdl-o-007, kdl-o-037. Belgilar: Material Icons (Apache 2.0).',
]
text(s, Inches(0.6), Inches(2.0), Inches(12.1), Inches(4.4), SOURCES, 15, INK, space=6)
NOTE_END = ('Taqdimot qoralama: ustoz mazmunni tekshirib, mahalliy tartiblarni (muassasa SOP, '
            'ekspozitsiyadan keyingi protokol, chiqindi sinflari) qo‘shishi kerak.')
notes(s, NOTE_END)
record('Manbalar', SOURCES, NOTE_END)

out_pptx = sys.argv[1] if len(sys.argv) > 1 else 'LabGuide_1-mavzu.pptx'
prs.save(out_pptx)
print(len(prs.slides), 'slayd')

# Ilovada slayd PPTX ko'rinishida — rasm (assets/lectures/d001/NN.jpg):
#   soffice --headless --convert-to pdf chiqish.pptx
#   pdftoppm -png -scale-to-x 1600 -scale-to-y 900 chiqish.pdf p  -> JPG (q82)
# Matn (sarlavha, bandlar) ekran o'quvchi va qidiruv uchun saqlanadi.
assert len(DATA) == len(prs.slides), (len(DATA), len(prs.slides))
app = [{**d, 'image': f'assets/lectures/d001/{i + 1:02d}.jpg'} for i, d in enumerate(DATA)]
out = os.path.join(REPO, 'assets', 'lectures', 'd001.json')
with open(out, 'w', encoding='utf-8') as f:
    json.dump(dict(
        schema=1, topics=['d001-P', 'd001-L'], lang='uz', status='draft',
        title='Laborator xizmatni tashkillashtirish. Tibbiy etika va deontologiya',
        slides=app), f, ensure_ascii=False, indent=1)
    f.write('\n')
print(out)
