"""Amaliy mashg'ulot protokollari to'plamining invariantlari.

  python3 -m unittest discover -s tool/slides -p 'test_*.py'

Tekshiriladi: ma'lumot to'liq (validate), uch tilda, har kontent slaydda manba (muallif + qisqa nom +
bo'lim/sahifa), oxirgi slayd "Adabiyotlar", qoralama yozuvi yo'q, protokollar o'quv kunlariga
BIRIKTIRILMAGAN (id raqamsiz, curriculum kunlari bilan kesishmaydi), render joylashtirish
ogohlantirishlarisiz o'tadi, hisobotda "kunma-kun mavzu yo'q" aytilgan.
"""
import os
import re
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import deck_content as dc  # noqa: E402
import practice_decks as pd  # noqa: E402

try:
    import pptx  # noqa: F401
    import make_topic_decks as mk
    HAVE_PPTX = True
except ImportError:  # pragma: no cover
    HAVE_PPTX = False

_repo = None


def repo():
    global _repo
    if _repo is None:
        _repo = dc.Repo()
    return _repo


class PracticeData(unittest.TestCase):
    def test_validate_is_clean(self):
        self.assertEqual(pd.validate(repo()), [])

    def test_there_is_a_set(self):
        self.assertGreaterEqual(len(pd.protocols(repo())), 20)

    def test_ids_are_unnumbered_and_do_not_collide_with_days(self):
        ids = [p['id'] for p in pd.protocols(repo())]
        self.assertEqual(len(ids), len(set(ids)))
        day_ids = {f'd{n:03d}' for n in repo().days}
        for pid in ids:
            self.assertRegex(pid, r'^[a-z0-9-]+$')
            self.assertNotRegex(pid, r'^d\d{3}')
            self.assertNotIn(pid, day_ids)

    def test_practice_is_kept_out_of_day_topics(self):
        # books_extra topics = faqat kunlar; amaliy protokollar kunga bog'lanmaydi
        ids = {p['id'] for p in pd.protocols(repo())}
        for key in repo().books['topics']:
            self.assertNotIn(key, ids)

    def test_no_draft_text_and_no_day_mapping_in_data(self):
        bad = re.compile(r'qoralama|mutaxassis tekshiruvidan|черновик|draft', re.I)
        for p in pd.protocols(repo()):
            blob = repr(p)
            self.assertIsNone(bad.search(blob), p['id'])
            self.assertIsNone(re.search(r'\bd0\d\d\b|\bd1[0-4]\d\b', blob), f'{p["id"]}: kun raqami tilga olingan')

    def test_every_content_slide_cites_book_and_location(self):
        for p in pd.protocols(repo()):
            for n, sl in enumerate(p['slides'], 1):
                if sl['type'] == 'gallery':
                    continue
                self.assertTrue(sl['refs'], f'{p["id"]}#{n}')
                for r in sl['refs']:
                    for lg in dc.LANGS:
                        self.assertTrue(dc.L(r['loc'], lg).strip(), f'{p["id"]}#{n}/{lg}')

    def test_every_protocol_has_safety_and_sources_cover_the_set(self):
        for p in pd.protocols(repo()):
            labels = [s.get('label') for s in p['slides']]
            self.assertIn('safety', labels, p['id'])
        self.assertGreaterEqual(len(repo().practice['books']), 3)

    def test_medical_audit_2026_10_17_unsafe_text_is_gone(self):
        """docs/AUDIT_2026-10-17_NEW_SECTIONS_MEDICAL.md: xavfli/eskirgan matn qaytmasin (uch tilda)."""
        bad = {
            'en': [r'cap is put on the needle', r'needle is detached', r're-used lancets',
                   r'single-use where possible', r'10–15 seconds', r'\bBlow the blood',
                   r'earlobe;', r'big toe', r'swab moistened with alcohol on the puncture',
                   r'within 45 minutes', r'later than 45 minutes', r'180–320',
                   r'Primary thrombocytosis is above 400', r'from 2 to 72 hours',
                   r'do not remove contact lenses'],
            'ru': [r'надевают колпачок, иглу отсоединяют', r'повторно используемые ланцеты',
                   r'10–15 секунд', r'[Вв]ыдувают', r'в течение 45 минут', r'180–320',
                   r'через 2–72 часа', r'линзы не снимайте'],
            'uz': [r'qalpoqchasi kiydiriladi', r'Qayta ishlatilgan lanset', r'10–15 soniya',
                   r'puflab', r'45 daqiqa ichida', r'180–320', r'2 soatdan 72 soatgacha',
                   r'linzani olmang'],
        }
        for p in pd.protocols(repo()):
            for n, sl in enumerate(p['slides'], 1):
                if sl['type'] != 'text':
                    continue
                for lg, pats in bad.items():
                    text = '\n'.join(sl['bullets'][lg])
                    for pat in pats:
                        self.assertIsNone(re.search(pat, text), f'{p["id"]}#{n}/{lg}: {pat}')

    def test_medical_audit_2026_10_17_who_rules_present(self):
        prots = {p['id']: p for p in pd.protocols(repo())}

        def text(pid, lg='en'):
            return '\n'.join(b for s in prots[pid]['slides'] if s['type'] == 'text' for b in s['bullets'][lg])

        bio = text('laboratoriya-biologik-xavfsizlik-sanepid')
        self.assertIn('Do not recap or detach the needle', bio)
        self.assertIn('sharps container', bio)
        self.assertIn('at least 30 seconds', bio)
        self.assertIn('as soon as possible', bio)
        self.assertIn('no later than 72 hours', bio)
        self.assertIn('28 days', bio)
        self.assertIn('kamida 30 soniya', text('laboratoriya-biologik-xavfsizlik-sanepid', 'uz'))
        self.assertIn('не менее 30 секунд', text('laboratoriya-biologik-xavfsizlik-sanepid', 'ru'))
        cap = text('kapillyar-qon-olish')
        self.assertIn('single-use lancets only', cap)
        self.assertIn('clean dry gauze', cap)
        self.assertIn('medial or lateral surface of the heel', cap)
        for pid in ('gemoglobin-gemiglobintsianid', 'eritrotsitlarni-sanash-kamera', 'leykotsitlarni-sanash-kamera'):
            self.assertIn('never by mouth', text(pid), pid)
        # Har bir WHO 2010 iqtiboslangan slaydda bo'lim/sahifa bor.
        who_refs = [r for p in pd.protocols(repo()) for s in p['slides'] for r in s.get('refs', [])
                    if r.get('book') == 'who-phlebotomy-2010']
        self.assertGreaterEqual(len(who_refs), 5)
        for r in who_refs:
            for lg in dc.LANGS:
                self.assertRegex(r['loc'][lg], r'\d', r)

    def test_report_states_there_is_no_day_mapping(self):
        # Manbada kun <-> texnika moslik yo'q; bog'lanish faqat LabGuide tavsiyasi sifatida (D-67).
        text = '\n'.join(pd.report_section(repo()))
        self.assertIn('kunma-kun mavzu ko‘rsatmaydi', text)
        self.assertIn('LabGuide tavsiya etgan amaliy reja', text)
        self.assertIn('manbadagi moslik emas', text)
        self.assertIn('Manbalarda bor, lekin olinmagan', text)
        self.assertNotIn('qoralama', text.lower())


class PracticeDecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.decks = []
        for p in pd.protocols(repo()):
            for lg in dc.LANGS:
                deck, cov = pd.build_practice_deck(repo(), p, lg)
                assert deck is not None, p['id']
                cls.decks.append(deck)

    def test_structure(self):
        for d in self.decks:
            kinds = [s.kind for s in d.slides]
            self.assertEqual(kinds[0], 'title')
            self.assertEqual(kinds[1], 'outline')
            self.assertEqual(kinds[-1], 'sources')
            self.assertEqual(d.slides[-1].title, dc.T('sources', d.lang))
            self.assertGreaterEqual(len(d.content_slides), 3, d.day.id)

    def test_sources_on_every_content_slide_and_listed(self):
        for d in self.decks:
            used = set()
            for i, s in enumerate(d.slides):
                if s.is_content:
                    self.assertTrue(s.sources, f'{d.day.id}/{d.lang} #{i} "{s.title}": manba yo‘q')
                used |= {r['key'] for r in s.sources}
            self.assertEqual(used - set(d.sources.entries), set())
            listed = d.slides[-1].data['items']
            for key in used:
                self.assertIn(d.sources.bib(key), listed, f'{d.day.id}/{d.lang} {key}')

    def test_books_show_their_locations_on_the_references_slide(self):
        for d in self.decks:
            joined = ' '.join(d.slides[-1].data['items'])
            self.assertRegex(joined, r'\d', d.day.id)

    def test_three_languages_really_differ(self):
        by = {}
        for d in self.decks:
            by.setdefault(d.day.id, {})[d.lang] = d
        for pid, langs in by.items():
            self.assertRegex(langs['ru'].slides[0].title, r'[А-Яа-я]', pid)
            self.assertNotRegex(langs['uz'].slides[0].title, r'[А-Яа-я]', pid)
            self.assertRegex(langs['en'].slides[0].title, r'[A-Za-z]', pid)
            self.assertEqual(len({langs[lg].slides[0].title for lg in dc.LANGS}), 3, pid)

    def test_no_draft_text_in_slides(self):
        bad = re.compile(r'qoralama|mutaxassis tekshiruvidan|черновик|draft', re.I)
        for d in self.decks:
            for s in d.slides:
                self.assertIsNone(bad.search(s.title + str(s.data)), d.day.id)


@unittest.skipUnless(HAVE_PPTX, 'python-pptx kerak')
class PracticeRender(unittest.TestCase):
    def test_all_decks_render_without_layout_warnings(self):
        from pptx import Presentation
        out = tempfile.mkdtemp()
        before = len(mk.WARNINGS)
        refs_title = {lg: dc.T('sources', lg) for lg in dc.LANGS}
        for p in pd.protocols(repo()):
            for lg in dc.LANGS:
                deck, _ = pd.build_practice_deck(repo(), p, lg)
                path = os.path.join(out, lg, f'{p["id"]}.pptx')
                n = mk.render_deck(deck, path)
                prs = Presentation(path)
                slides = list(prs.slides)
                self.assertEqual(len(slides), n)
                self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=2)
                for i, slide in enumerate(slides, 1):
                    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
                    self.assertIn('LabGuide', texts, f'{p["id"]}/{lg} #{i}')
                    self.assertIn(f'{i} / {n}', texts)
                    self.assertNotIn('qoralama', ' '.join(texts).lower())
                last = [sh.text_frame.text for sh in slides[-1].shapes if sh.has_text_frame]
                self.assertIn(refs_title[lg], last, f'{p["id"]}/{lg}')
        self.assertEqual(mk.WARNINGS[before:], [], 'joylashtirish ogohlantirishlari')


if __name__ == '__main__':
    unittest.main()
