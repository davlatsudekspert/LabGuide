"""Mavzu taqdimotlari generatorining invariantlari.

  python3 -m unittest discover -s tool/slides -p 'test_*.py'

Tekshiriladi: Dart'dan ko'chirilgan JSON mos; har kontent slaydida manba bor;
bo'sh slayd yo'q; toifa testlaridan faqat kaliti aniq savollar; footer faqat
"LabGuide" (qoralama yozuvi yo'q); oxirgi slayd — "Adabiyotlar"; matn
joylashtirishda sig'maslik ogohlantirishi yo'q.
"""
import os
import re
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import deck_content as dc  # noqa: E402
import extract_dart_content as ex  # noqa: E402

try:
    import pptx  # noqa: F401
    import make_topic_decks as mk
    HAVE_PPTX = True
except ImportError:  # pragma: no cover
    HAVE_PPTX = False

REPO = dc.REPO
_repo = None


def repo():
    global _repo
    if _repo is None:
        _repo = dc.Repo()
    return _repo


class DartSync(unittest.TestCase):
    def test_json_matches_dart(self):
        cur = ex._read(ex.OUT)
        self.assertEqual(cur, ex.dump(ex.build()),
                         'app_content.json Dart bilan mos emas: python3 tool/slides/extract_dart_content.py')

    def test_counts_match_dart_sources(self):
        def n(rel, cls):
            src = ex._read(os.path.join(REPO, rel))
            # `ClassName(` chaqiriqlari, e'lon (class X { const X( ...) ) va konstruktor emas
            return len(re.findall(r'^\s{2}' + cls + r'\(', src, re.M))
        app = repo().app
        self.assertEqual(len(app['cellGuides']),
                         n('lib/features/differential/differential_content.dart', 'CellGuide'))
        self.assertEqual(len(app['confusions']),
                         n('lib/features/differential/differential_content.dart', 'Confusion'))
        self.assertEqual(len(app['diffFindings']),
                         n('lib/features/differential/differential_content.dart', 'DiffFinding'))
        self.assertEqual(len(app['cytoCells']),
                         n('lib/features/cytology/cytology_content.dart', 'CytoCell'))
        self.assertEqual(len(app['bethesdaItems']),
                         n('lib/features/cytology/cytology_content.dart', 'BethesdaItem'))

    def test_three_languages_everywhere(self):
        app = repo().app

        def walk(o, path=''):
            if isinstance(o, dict):
                if set(o) >= {'uz', 'ru', 'en'}:
                    for lg in dc.LANGS:
                        self.assertTrue(str(o[lg]).strip(), f'{path}.{lg} bo‘sh')
                else:
                    for k, v in o.items():
                        walk(v, f'{path}.{k}')
            elif isinstance(o, list):
                for i, v in enumerate(o):
                    walk(v, f'{path}[{i}]')
        for k in ('cellGuides', 'confusions', 'diffFindings', 'cytoCells', 'bethesdaItems', 'refTopics'):
            walk(app[k], k)

    def test_images_exist(self):
        r = repo()
        for c in r.app['cellGuides']:
            self.assertTrue(os.path.exists(os.path.join(REPO, 'assets/differential', c['id'] + '.png')), c['id'])
        for c in r.app['cytoCells']:
            self.assertTrue(os.path.exists(os.path.join(REPO, 'assets/cytology', c['id'] + '.png')), c['id'])
        for i in r.atlas['images']:
            self.assertTrue(os.path.exists(os.path.join(REPO, i['asset'])), i['id'])


class BooksExtra(unittest.TestCase):
    def test_entries_complete_in_three_languages_and_cited(self):
        r = repo()
        self.assertTrue(r.books['topics'])
        for day_id, entries in r.books['topics'].items():
            self.assertIn(int(day_id[1:]), r.days, day_id)
            for e in entries:
                for lg in dc.LANGS:
                    self.assertTrue(dc.L(e['title'], lg).strip())
                    if 'table' in e:
                        self.assertTrue(e['table']['rows'][lg] and e['table']['header'][lg])
                        for row in e['table']['rows'][lg]:
                            self.assertEqual(len(row), len(e['table']['header'][lg]))
                    else:
                        self.assertTrue(e['bullets'][lg], f'{day_id}/{lg}')
                self.assertTrue(e['refs'], f'{day_id}: manba yo‘q')
                for ref in e['refs']:
                    self.assertEqual(r.src[ref['pack']]['kind'], 'book')
                    for lg in dc.LANGS:
                        self.assertTrue(dc.L(ref['loc'], lg).strip(), f'{day_id}: bo‘lim/sahifa yo‘q')

    def test_book_slides_are_in_decks_and_listed_with_locations(self):
        r = repo()
        for day_id in r.books['topics']:
            day = r.days[int(day_id[1:])]
            for lang in dc.LANGS:
                deck, _ = dc.build_deck(r, day, lang)
                book = [s for s in deck.slides if s.eyebrow == dc.T('book_extra', lang)]
                self.assertEqual(len(book), len(r.books['topics'][day_id]))
                lst = ' '.join(deck.slides[-1].data['items'])
                for e in r.books['topics'][day_id]:
                    for ref in e['refs']:
                        self.assertIn(dc.L(ref['loc'], lang), lst)


def _slide_is_empty(s):
    d = s.data
    if s.kind in ('text',):
        return not any(b['text'].strip() for b in d['blocks'])
    if s.kind == 'columns':
        return not (d['single'] or any(c['blocks'] for c in d['cols']))
    if s.kind == 'table':
        return not d['rows']
    if s.kind == 'card':
        return not d['rows'] or not s.image
    if s.kind == 'pair':
        return not (d['a']['image'] and d['b']['image'])
    if s.kind == 'gallery':
        return not d['items']
    if s.kind == 'qa_q':
        return not (d['question'].strip() and len(d['options']) >= 2)
    if s.kind == 'qa_a':
        return not d['options'][d['correct']].strip()
    if s.kind in ('outline', 'sources'):
        return not d['items']
    return False


class DeckInvariants(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.decks = []
        for day in repo().days.values():
            for lang in dc.LANGS:
                deck, cov = dc.build_deck(repo(), day, lang)
                if deck:
                    cls.decks.append(deck)

    def test_there_are_decks(self):
        self.assertGreater(len({d.day.id for d in self.decks}), 40)

    def test_sources_on_every_content_slide(self):
        for d in self.decks:
            for i, s in enumerate(d.slides):
                if s.is_content:
                    self.assertTrue(s.sources, f'{d.day.id}/{d.lang} #{i} {s.kind} "{s.title}": manba yo‘q')
                    for r in s.sources:
                        self.assertIn(r['key'], d.sources.entries)

    def test_no_empty_slides(self):
        for d in self.decks:
            for i, s in enumerate(d.slides):
                self.assertFalse(_slide_is_empty(s), f'{d.day.id}/{d.lang} #{i} {s.kind} "{s.title}" bo‘sh')
                self.assertTrue(s.title.strip(), f'{d.day.id}/{d.lang} #{i}: sarlavha yo‘q')

    def test_structure_title_outline_sources(self):
        for d in self.decks:
            kinds = [s.kind for s in d.slides]
            self.assertEqual(kinds[0], 'title')
            self.assertEqual(kinds[1], 'outline')
            self.assertEqual(kinds[-1], 'sources')
            self.assertEqual(d.slides[-1].title, dc.T('sources', d.lang))
            self.assertGreaterEqual(len(d.content_slides), 1)

    def test_sources_slide_lists_every_used_source(self):
        for d in self.decks:
            used = {r['key'] for s in d.slides for r in s.sources}
            self.assertEqual(used - set(d.sources.entries), set())
            listed = d.slides[-1].data['items']
            for key in used:
                self.assertIn(d.sources.bib(key), listed, f'{d.day.id}/{d.lang} {key}')

    def test_only_verified_toifa_tests(self):
        r = repo()
        for day in r.days.values():
            for q in dc.test_questions(r, day):
                self.assertEqual(q['verdict'], 'ok')
                self.assertEqual(len(q['key']), 1)
                self.assertFalse(q.get('unverified'))
                self.assertFalse(q.get('held'))

    def test_quiz_slides_always_have_sources(self):
        for d in self.decks:
            for s in d.slides:
                if s.kind in ('qa_q', 'qa_a'):
                    self.assertTrue(s.sources)

    def test_atlas_credit_present(self):
        for d in self.decks:
            for s in d.slides:
                if s.kind == 'gallery':
                    for it in s.data['items']:
                        self.assertIn(dc.T('photo', d.lang), it['credit'])
                        self.assertRegex(it['credit'], r'CC|Public|PD|CDC')

    def test_languages_differ(self):
        by = {}
        for d in self.decks:
            by.setdefault(d.day.id, {})[d.lang] = d
        d012 = by['d012']
        self.assertNotEqual(d012['uz'].slides[0].title, d012['ru'].slides[0].title)
        self.assertRegex(d012['ru'].slides[0].title, r'[А-Яа-я]')
        self.assertRegex(d012['en'].slides[0].title, r'[A-Za-z]')

    def test_no_draft_footer_text_in_content(self):
        bad = re.compile(r'qoralama|mutaxassis tekshiruvidan|черновик|draft', re.I)
        for d in self.decks:
            for s in d.slides:
                if s.kind in ('title', 'outline'):
                    self.assertIsNone(bad.search(s.title + str(s.data)))

    def test_pick_days(self):
        r = repo()
        self.assertEqual([x.id for x in mk.pick_days(r, ['d012-P', '12', 'D012-L'], False)] if HAVE_PPTX else ['d012'],
                         ['d012'])


@unittest.skipUnless(HAVE_PPTX, 'python-pptx kerak')
class Render(unittest.TestCase):
    def _render(self, day_id, lang):
        r = repo()
        day = r.days[int(day_id[1:])]
        deck, _ = dc.build_deck(r, day, lang)
        self.assertIsNotNone(deck)
        path = os.path.join(tempfile.mkdtemp(), f'{day_id}_{lang}.pptx')
        before = len(mk.WARNINGS)
        n = mk.render_deck(deck, path)
        self.assertEqual(mk.WARNINGS[before:], [], 'joylashtirish ogohlantirishlari')
        from pptx import Presentation
        return Presentation(path), deck, n

    def test_footer_labguide_only_and_widescreen(self):
        for day_id, lang in (('d012', 'uz'), ('d007', 'ru'), ('d035', 'en')):
            prs, deck, n = self._render(day_id, lang)
            self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=2)
            self.assertGreaterEqual(n, len(deck.slides))
            for i, slide in enumerate(prs.slides, 1):
                texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
                self.assertIn('LabGuide', texts, f'{day_id}/{lang} #{i}: footer yo‘q')
                self.assertIn(f'{i} / {n}', texts)
                joined = ' '.join(texts).lower()
                self.assertNotIn('qoralama', joined)
                self.assertNotIn('mutaxassis tekshiruvidan', joined)

    def test_every_content_page_has_source_line_and_last_is_references(self):
        prs, deck, n = self._render('d012', 'uz')
        slides = list(prs.slides)
        first_prefix = dc.T('source_prefix', 'uz') + ':'
        refs_title = dc.T('sources', 'uz')
        src_pages = 0
        for i, slide in enumerate(slides):
            texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
            if i in (0, 1) or refs_title in texts:
                continue
            if any(t.startswith(first_prefix) for t in texts):
                src_pages += 1
        # sarlavha, reja va "Adabiyotlar" sahifalaridan boshqa hammasida manba qatori bor
        ref_pages = sum(1 for sl in slides
                        if refs_title in [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame])
        self.assertGreaterEqual(ref_pages, 1)
        self.assertEqual(src_pages, len(slides) - 2 - ref_pages)
        last = [sh.text_frame.text for sh in slides[-1].shapes if sh.has_text_frame]
        self.assertIn(refs_title, last)

    def test_all_images_embedded_have_alt_text(self):
        prs, deck, n = self._render('d035', 'uz')
        for slide in prs.slides:
            for sh in slide.shapes:
                if sh.shape_type == 13:  # PICTURE
                    self.assertTrue(sh._element.nvPicPr.cNvPr.get('descr') is not None)


if __name__ == '__main__':
    unittest.main()
