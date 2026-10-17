#!/usr/bin/env python3
"""Curriculum mavzulari uchun PPTX taqdimotlar (16:9, televizor uchun).

Mazmun faqat ilovadagi tekshirilgan kontentdan yig'iladi (deck_content.py):
analit/holat kartalari, leykoformula va sitologiya sxemalari, mikroskopik atlas
fotolari, toifa savollari, "Jadvallar"; ixtiyoriy — ustoz kitoblaridan o'z
so'zim bilan qisqa mazmun (data/books_extra.json). Kontent bo'lmasa slayd yo'q.
Har kontent slaydida qisqa manba qatori, oxirgi slayd — "Adabiyotlar".

Ishlatish (PPTX'lar repoga qo'shilmaydi — --out ni repodan tashqariga bering):
  python3 tool/slides/make_topic_decks.py --all --lang uz --out /tmp/decks
  python3 tool/slides/make_topic_decks.py --topic d012 --lang uz,ru,en --out /tmp/decks
  python3 tool/slides/make_topic_decks.py --all --dry-run --report docs/DECKS_COVERAGE.md

Kerak: python-pptx, Pillow (venv: python3 -m venv /tmp/venv && /tmp/venv/bin/pip install python-pptx).
Shrift: Inter (assets/fonts); kompyuterda bo'lmasa PowerPoint o'rnini bosadi.
"""
import argparse
import math
import os
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import deck_content as dc  # noqa: E402

REPO = dc.REPO

# ------------------------------------------------------------- ranglar ------
# lib/design/tokens.dart (LgPalette.light) + qorong'i sarlavha foni
BG = 'F3F3EC'
PAPER = 'FFFEF9'
INK = '183A36'
SUB = '536C65'
LINE = 'DCE4D9'
BRAND = '194C40'
SOFT = 'E8EEDB'
AMBER = '7B4D0B'
AMBER_BG = 'FFF1DB'
DEEP = '122D2A'
LIME = 'C8E27A'
WHITE = 'FFFFFF'
PALE = 'C9D8C2'

FONT = 'Inter'
W_IN, H_IN = 13.333, 7.5
MX = 0.75                       # chap/o'ng chekka
CW = W_IN - 2 * MX              # kontent kengligi
BODY_BOTTOM = 6.32
SRC_Y = 6.42
FOOT_Y = 7.04
LH = 1.25                       # qator balandligi koeffitsienti (xavfsiz)
SAFETY = 1.07                   # o'lchangan matn kengligiga zaxira

_pptx = {}
WARNINGS = []   # joylashtirib bo'lmagan holatlar (render paytida to'planadi)


def _imports():
    """python-pptx faqat render paytida kerak (test mazmunni PPTX'siz tekshiradi)."""
    if _pptx:
        return _pptx
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
    from lxml import etree
    _pptx.update(locals())
    return _pptx


# ------------------------------------------------------------ o'lchash ------

_fonts = {}


def _pil_font(bold):
    from PIL import ImageFont
    if bold not in _fonts:
        _fonts[bold] = ImageFont.truetype(
            os.path.join(REPO, 'assets', 'fonts', 'Inter-700.ttf' if bold else 'Inter-400.ttf'), 200)
    return _fonts[bold]


_wcache = {}


def text_w(s, size, bold=False):
    """Matn kengligi (pt)."""
    k = (s, bold)
    w = _wcache.get(k)
    if w is None:
        w = _wcache[k] = _pil_font(bold).getlength(s)
    return w * size / 200.0


def count_lines(text, size, width_in, bold=False, indent_in=0.0):
    """So'zma-so'z qatorlarga bo'lish (so'z kengliklari keshlanadi)."""
    avail = (width_in - indent_in) * 72.0 / SAFETY
    space = text_w(' ', size, bold)
    total = 0
    for para in str(text).split('\n'):
        n, cur = 1, 0.0
        for w in para.split(' '):
            ww = text_w(w, size, bold)
            if cur == 0.0:
                cur = ww
            elif cur + space + ww <= avail:
                cur += space + ww
            else:
                n += 1
                cur = ww
            if ww > avail:      # juda uzun so'z (URL): qo'shimcha qatorlar
                extra = int(ww // avail)
                n += extra
                cur = ww - extra * avail
        total += n
    return total


class Para:
    def __init__(self, text, size, bold=False, italic=False, color=INK, after=0.0,
                 bullet=None, indent=0.0, align='l', label=None):
        self.text, self.size, self.bold, self.italic, self.color = text, size, bold, italic, color
        self.after, self.bullet, self.indent, self.align = after, bullet, indent, align
        self.label = label   # satr boshidagi qalin yorliq (bir qatorda)

    def height_in(self, width_in):
        full = f'{self.label}: {self.text}' if self.label else self.text
        # qalin yorliq biroz kengroq: shriftni 3% kattaroq o'lchaymiz
        n = count_lines(full, self.size * (1.03 if self.label else 1.0), width_in, self.bold, self.indent)
        return (n * self.size * LH + self.after) / 72.0


def paras_height(paras, width_in):
    return sum(p.height_in(width_in) for p in paras)


# ------------------------------------------------------------ ibtidoiylar ---

def rgb(h):
    return _imports()['RGBColor'].from_string(h)


def I(x):  # noqa: E743
    return _imports()['Inches'](x)


def set_bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = rgb(color)


def add_rect(slide, x, y, w, h, fill=None, line=None, radius=None, line_w=0.75):
    p = _imports()
    kind = p['MSO_SHAPE'].ROUNDED_RECTANGLE if radius else p['MSO_SHAPE'].RECTANGLE
    sh = slide.shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    if radius:
        sh.adjustments[0] = min(0.5, radius / max(min(w, h), 0.01))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = p['Pt'](line_w)
    else:
        sh.line.fill.background()
    st = sh._element.find(_imports()['qn']('p:style'))
    if st is not None:
        sh._element.remove(st)
    return sh


def _set_bullet(par, kind, color, indent_in):
    p = _imports()
    qn, etree = p['qn'], p['etree']
    pPr = par._p.get_or_add_pPr()
    emu = int(indent_in * 914400)
    pPr.set('marL', str(emu))
    pPr.set('indent', str(-emu))
    clr = etree.SubElement(pPr, qn('a:buClr'))
    srgb = etree.SubElement(clr, qn('a:srgbClr'))
    srgb.set('val', color)
    if kind == 'num':
        el = etree.SubElement(pPr, qn('a:buAutoNum'))
        el.set('type', 'arabicPeriod')
    else:
        el = etree.SubElement(pPr, qn('a:buChar'))
        el.set('char', '•')


def add_text(slide, x, y, w, h, paras, anchor='t', inset=0.0, autoshrink=False):
    p = _imports()
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = I(inset)
    tf.margin_top = tf.margin_bottom = I(0.02)
    tf.vertical_anchor = {'t': p['MSO_ANCHOR'].TOP, 'm': p['MSO_ANCHOR'].MIDDLE,
                          'b': p['MSO_ANCHOR'].BOTTOM}[anchor]
    first = True
    for pa in paras:
        par = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        par.alignment = {'l': p['PP_ALIGN'].LEFT, 'c': p['PP_ALIGN'].CENTER,
                         'r': p['PP_ALIGN'].RIGHT}[pa.align]
        par.line_spacing = 1.0
        par.space_after = p['Pt'](pa.after)
        par.space_before = p['Pt'](0)
        if pa.label:
            lr = par.add_run()
            lr.text = pa.label + ': '
            lf = lr.font
            lf.name = FONT
            lf.size = p['Pt'](pa.size)
            lf.bold = True
            lf.color.rgb = rgb(BRAND)
        run = par.add_run()
        run.text = pa.text
        f = run.font
        f.name = FONT
        f.size = p['Pt'](pa.size)
        f.bold = pa.bold
        f.italic = pa.italic
        f.color.rgb = rgb(pa.color)
        if pa.bullet:
            _set_bullet(par, pa.bullet, BRAND, pa.indent)
    return tb


def add_image(slide, rel, x, y, w, h, descr='', align='c'):
    """Rasmni (x,y,w,h) qutiga nisbatini saqlab joylaydi."""
    from PIL import Image
    path = os.path.join(REPO, rel)
    with Image.open(path) as im:
        iw, ih = im.size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    px = x + (w - pw) / 2 if align == 'c' else x
    py = y + (h - ph) / 2
    pic = slide.shapes.add_picture(path, I(px), I(py), I(pw), I(ph))
    pic._element.nvPicPr.cNvPr.set('descr', descr)
    return pic, (px, py, pw, ph)


# ------------------------------------------------------------- umumiy qism --

class Ctx:
    def __init__(self, deck):
        p = _imports()
        self.deck = deck
        self.lang = deck.lang
        self.prs = p['Presentation']()
        self.prs.slide_width = I(W_IN)
        self.prs.slide_height = I(H_IN)
        self.blank = self.prs.slide_layouts[6]
        self.pages = []   # (slide, spec, dark)

    def new(self, spec, dark=False):
        s = self.prs.slides.add_slide(self.blank)
        set_bg(s, DEEP if dark else BG)
        self.pages.append((s, spec, dark))
        return s


def source_line(ctx, spec):
    """Qisqa manba qatori (2 qatorga sig'adigan)."""
    refs = spec.sources
    if not refs:
        return ''
    src = ctx.deck.sources
    cites, seen = [], set()
    for r in refs:
        c = src.short_cite(r)
        if c not in seen:
            seen.add(c)
            cites.append(c)
    prefix = dc.T('source_prefix', ctx.lang) + ': '
    for k in range(min(len(cites), 4), 0, -1):
        s = prefix + '; '.join(cites[:k])
        if k < len(cites):
            s += ' · ' + dc.T('more_sources', ctx.lang, n=len(cites) - k)
        if count_lines(s, 13, CW) <= 2:
            return s
    s = prefix + dc.T('more_sources', ctx.lang, n=len(cites))
    return s


def chrome(ctx, slide, spec, n, total, dark=False):
    if not dark and spec.is_content:
        line = source_line(ctx, spec)
        if line:
            add_text(slide, MX, SRC_Y, CW, 0.5, [Para(line, 13, color=SUB)])
    fg = PALE if dark else SUB
    logo = os.path.join(REPO, 'assets', 'images', 'logo_mark.png')
    slide.shapes.add_picture(logo, I(MX), I(FOOT_Y), I(0.3), I(0.3)).\
        _element.nvPicPr.cNvPr.set('descr', 'LabGuide')
    add_text(slide, MX + 0.4, FOOT_Y + 0.02, 3, 0.3, [Para('LabGuide', 13, bold=True, color=fg)], anchor='m')
    add_text(slide, W_IN - MX - 1.5, FOOT_Y + 0.02, 1.5, 0.3,
             [Para(f'{n} / {total}', 13, color=fg, align='r')], anchor='m')


def title_block(ctx, slide, title, eyebrow, page=None):
    """Sarlavha bloki; kontent boshlanadigan y ni qaytaradi."""
    pg = f'{page[0]}/{page[1]}' if page and page[1] > 1 else ''
    ey = f'{eyebrow} · {pg}' if (eyebrow and pg) else (eyebrow or pg)
    y = 0.36
    if ey:
        add_text(slide, MX, y, CW, 0.3, [Para(ey.upper() if ctx.lang != 'ru' else ey, 14, bold=True, color=BRAND)])
    size = 34
    while size > 24 and count_lines(title, size, CW, True) > 2:
        size -= 2
    n = min(2, count_lines(title, size, CW, True))
    th = n * size * LH / 72.0
    add_text(slide, MX, y + 0.32, CW, th + 0.05, [Para(title, size, bold=True, color=INK)])
    return y + 0.32 + th + 0.2


# ---- matn bloklari (paragraflar) ---------------------------------------------

def block_paras(blocks, size, opts):
    bullets = opts.get('bullets')
    numbered = opts.get('numbered')
    inline = opts.get('inline')
    out = []
    for b in blocks:
        if inline and b.get('label'):
            out.append(Para(b['text'], size, label=b['label'], after=size * 0.45))
            continue
        if b.get('label'):
            out.append(Para(b['label'], max(size - 4, 16), bold=True, color=BRAND, after=2))
        bl = 'num' if numbered else ('dot' if (bullets or b.get('bullet')) else None)
        out.append(Para(b['text'], size, bullet=bl, indent=0.42 if numbered else (0.32 if bl else 0.0),
                        after=size * 0.55))
    return out


def split_sentences(text, seps='.!?'):
    parts = re.split(r'(?<=[' + re.escape(seps) + r'])\s+', text)
    return [p for p in parts if p]


def _paginate_at(blocks, width, height, opts, sz):
    """Berilgan shriftda sahifalarga bo'lish (bloklar kerak bo'lsa jumlalar bo'yicha)."""
    def fits(bs):
        return paras_height(block_paras(bs, sz, opts), width) <= height

    pages, cur, queue = [], [], list(blocks)
    while queue:
        b = queue.pop(0)
        if fits(cur + [b]):
            cur.append(b)
            continue
        if cur:
            pages.append(cur)
            cur = []
        if fits([b]):
            cur = [b]
            continue
        sents = split_sentences(b['text'])
        if any(not fits([{**b, 'text': x}]) for x in sents):
            sents = split_sentences(b['text'], '.!?;')
        if any(not fits([{**b, 'text': x}]) for x in sents):
            sents = split_sentences(b['text'], '.!?;:,')
        if len(sents) <= 1:
            cur = [b]   # bo'linmaydi (test sig'maslikni ushlaydi)
            continue
        piece, rest = [], []
        for i, s_ in enumerate(sents):
            if fits([{**b, 'text': ' '.join(piece + [s_])}]):
                piece.append(s_)
            else:
                rest = sents[i:]
                break
        if not piece:
            piece, rest = [sents[0]], sents[1:]
        pages.append([{**b, 'text': ' '.join(piece)}])
        if rest:
            queue.insert(0, {**b, 'label': None, 'text': ' '.join(rest)})
    if cur:
        pages.append(cur)
    return pages


def paginate_blocks(blocks, width, height, opts, sizes=(28, 26, 24, 22, 20), page_size=20, warn=True):
    """-> [(bloklar, shrift)] sahifalar. Bitta sahifaga sig'sa — eng katta mos
    shrift; bo'lmasa — hamma sahifada bir xil shrift (sahifalar soni kamayadi)."""
    for sz in sizes:
        if paras_height(block_paras(blocks, sz, opts), width) <= height:
            return [(blocks, sz)]
    base = _paginate_at(blocks, width, height, opts, page_size)
    best = page_size
    for sz in sizes:
        if sz > page_size and len(_paginate_at(blocks, width, height, opts, sz)) == len(base):
            best = sz
            break
    pages = _paginate_at(blocks, width, height, opts, best)
    for pg in pages:
        h = paras_height(block_paras(pg, best, opts), width)
        if warn and h > height + 0.03:
            WARNINGS.append(f'matn sig\'madi ({h:.2f} > {height:.2f} in): {pg[0]["text"][:50]}')
    return [(pg, best) for pg in pages]


# ------------------------------------------------------------ slayd turlari --

def r_title(ctx, spec):
    s = ctx.new(spec, dark=True)
    add_rect(s, MX, 1.55, 1.2, 0.09, fill=LIME)
    add_text(s, MX, 1.8, CW, 0.4, [Para(spec.eyebrow.upper() if ctx.lang != 'ru' else spec.eyebrow, 18, bold=True, color=LIME)])
    size = 46
    while size > 28 and count_lines(spec.title, size, CW - 1, True) * size * LH / 72 > 3.0:
        size -= 2
    th = count_lines(spec.title, size, CW - 1, True) * size * LH / 72
    add_text(s, MX, 2.35, CW - 1, th + 0.05, [Para(spec.title, size, bold=True, color=WHITE)])
    lines = [Para(spec.data['module'], 20, color=PALE, after=8)]
    lines += [Para(x, 18, color=PALE, after=4) for x in spec.data['lines']]
    add_text(s, MX, 2.35 + th + 0.35, CW - 1, 1.7, lines)


def r_outline(ctx, spec):
    items = [{'label': None, 'text': t} for t in spec.data['items']]
    pages = paginate_blocks(items, CW, BODY_BOTTOM - 1.5, {'bullets': True}, sizes=(26, 24, 22, 20))
    for i, (bl, sz) in enumerate(pages, 1):
        s = ctx.new(spec)
        y = title_block(ctx, s, spec.title, '', (i, len(pages)))
        add_text(s, MX, y + 0.1, CW, BODY_BOTTOM - y, block_paras(bl, sz, {'bullets': True}))


def _chips_height(chips, width):
    if not chips:
        return 0
    return count_lines('   ·   '.join(chips), 16, width) * 16 * LH / 72 + 0.15


def r_text(ctx, spec):
    d = spec.data
    opts = {'bullets': d.get('bullets'), 'numbered': d.get('numbered')}
    blocks = d['blocks']
    title_n = min(2, count_lines(spec.title, 34, CW, True))
    top = 0.36 + 0.32 + title_n * 34 * LH / 72.0 + 0.2
    chips = d.get('chips') or []
    lead = d.get('lead')
    extra = _chips_height(chips, CW)
    if lead:
        extra += count_lines(lead, 20, CW) * 20 * LH / 72 + 0.2
    avail = BODY_BOTTOM - top - extra
    pages = paginate_blocks(blocks, CW, avail, opts)
    for i, (bl, sz) in enumerate(pages, 1):
        s = ctx.new(spec)
        y = title_block(ctx, s, spec.title, spec.eyebrow, (i, len(pages)))
        if lead:
            h = count_lines(lead, 20, CW) * 20 * LH / 72
            add_text(s, MX, y, CW, h + 0.05, [Para(lead, 20, italic=True, color=SUB)])
            y += h + 0.2
        if chips:
            ch = count_lines('   ·   '.join(chips), 16, CW) * 16 * LH / 72
            add_text(s, MX, y, CW, ch + 0.05, [Para('   ·   '.join(chips), 16, color=SUB)])
            y += ch + 0.15
        if d.get('warn'):
            add_rect(s, MX - 0.15, y - 0.05, CW + 0.3, BODY_BOTTOM - y + 0.1, fill=AMBER_BG, radius=0.12)
        add_text(s, MX, y, CW, BODY_BOTTOM - y, block_paras(bl, sz, opts))


def r_columns(ctx, spec):
    d = spec.data
    cols = d['cols']
    if len(cols) == 2 and not d['single']:
        gap = 0.4
        cw = (CW - gap) / 2
        title_n = min(2, count_lines(spec.title, 34, CW, True))
        top = 0.36 + 0.32 + title_n * 34 * LH / 72.0 + 0.2
        head_h = 0.55
        avail = BODY_BOTTOM - top - head_h - 0.1
        for sz in (24, 22, 20, 18):
            ok = all(paras_height(block_paras(c['blocks'], sz, {'bullets': len(c['blocks']) > 1}), cw - 0.3) <= avail - 0.3
                     for c in cols)
            if ok:
                s = ctx.new(spec)
                y = title_block(ctx, s, spec.title, spec.eyebrow)
                for k, c in enumerate(cols):
                    x = MX + k * (cw + gap)
                    tone = BRAND if k == 0 else AMBER
                    bgc = SOFT if k == 0 else AMBER_BG
                    add_rect(s, x, y, cw, head_h, fill=tone, radius=0.1)
                    add_text(s, x + 0.2, y, cw - 0.4, head_h, [Para(c['head'], 22, bold=True, color=WHITE)], anchor='m')
                    add_rect(s, x, y + head_h + 0.05, cw, BODY_BOTTOM - y - head_h - 0.05, fill=bgc, radius=0.1)
                    add_text(s, x + 0.15, y + head_h + 0.15, cw - 0.3, BODY_BOTTOM - y - head_h - 0.3,
                             block_paras(c['blocks'], sz, {'bullets': len(c['blocks']) > 1}))
                return
    # ketma-ket: har ustun alohida (to'liq kenglik), sarlavha — yorliq sifatida
    blocks = []
    for c in cols:
        for i, b in enumerate(c['blocks']):
            blocks.append({'label': c['head'] if i == 0 else None, 'text': b['text']})
    blocks += d['single']
    sub = dc.Slide('text', title=spec.title, eyebrow=spec.eyebrow,
                   data={'blocks': blocks, 'chips': [], 'bullets': False}, sources=spec.sources)
    r_text(ctx, sub)


def table_layout(header, rows, widths, size, total_w):
    cw = [w * total_w for w in widths]
    pad = 0.09
    heights = []
    hh = max(count_lines(h, size, cw[i] - 2 * pad, True) for i, h in enumerate(header)) \
        if any(header) else 0
    hh = hh * size * LH / 72 + 2 * 0.06 if hh else 0
    for r in rows:
        n = max(count_lines(c, size, cw[i] - 2 * pad) for i, c in enumerate(r))
        heights.append(n * size * LH / 72 + 2 * 0.06)
    return cw, hh, heights


def r_table(ctx, spec):
    d = spec.data
    header, rows, widths = d['header'], d['rows'], d['widths']
    note = d.get('note') or ''
    title_n = min(2, count_lines(spec.title, 34, CW, True))
    top = 0.36 + 0.32 + title_n * 34 * LH / 72.0 + 0.2
    note_h = (count_lines(note, 14, CW) * 14 * LH / 72 + 0.12) if note else 0
    avail = BODY_BOTTOM - top - note_h
    chosen = None
    for sz in (22, 20, 18, 16):
        cw, hh, hs = table_layout(header, rows, widths, sz, CW)
        if hh + sum(hs) <= avail:
            chosen = [(rows, sz)]
            break
    if not chosen:
        # sahifalash: sahifalar soni eng kam bo'ladigan shrift (teng bo'lsa kattaroq)
        best = None
        for sz in (18, 16, 14):
            cw, hh, hs = table_layout(header, rows, widths, sz, CW)
            pages, cur, used = [], [], hh
            for r, h in zip(rows, hs):
                if cur and used + h > avail:
                    pages.append(cur)
                    cur, used = [], hh
                cur.append(r)
                used += h
            if cur:
                pages.append(cur)
            fits = all(hh + h <= avail for h in hs)
            key = (not fits, len(pages), -sz)
            if best is None or key < best[0]:
                best = (key, sz, pages, fits)
        _, sz, pages, fits = best
        if not fits:
            WARNINGS.append(f'jadval qatori sig\'madi: {spec.title[:50]}')
        chosen = [(pg, sz) for pg in pages]
    p = _imports()
    for i, (pg, sz) in enumerate(chosen, 1):
        s = ctx.new(spec)
        y = title_block(ctx, s, spec.title, spec.eyebrow, (i, len(chosen)))
        cw, hh, hs = table_layout(header, pg, widths, sz, CW)
        nrows = len(pg) + (1 if hh else 0)
        gf = s.shapes.add_table(nrows, len(header), I(MX), I(y), I(CW), I(hh + sum(hs)))
        tbl = gf.table
        tbl.first_row = bool(hh)
        tbl.horz_banding = False
        for j, w in enumerate(cw):
            tbl.columns[j].width = I(w)
        ri = 0
        if hh:
            tbl.rows[0].height = I(hh)
            for j, h in enumerate(header):
                _cell(tbl.cell(0, j), h, sz, True, WHITE, BRAND)
            ri = 1
        for k, r in enumerate(pg):
            tbl.rows[ri + k].height = I(hs[k])
            for j, c in enumerate(r):
                _cell(tbl.cell(ri + k, j), c, sz, j == 0, INK, SOFT if k % 2 else PAPER)
        if note:
            add_text(s, MX, y + hh + sum(hs) + 0.1, CW, note_h, [Para(note, 14, italic=True, color=SUB)])


def _cell(cell, text, size, bold, color, fill):
    p = _imports()
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    cell.margin_left = cell.margin_right = I(0.09)
    cell.margin_top = cell.margin_bottom = I(0.06)
    cell.vertical_anchor = p['MSO_ANCHOR'].TOP
    tf = cell.text_frame
    tf.word_wrap = True
    par = tf.paragraphs[0]
    par.line_spacing = 1.0
    run = par.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = p['Pt'](size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def r_card(ctx, spec):
    d = spec.data
    title_n = min(2, count_lines(spec.title, 34, CW, True))
    top = 0.36 + 0.32 + title_n * 34 * LH / 72.0 + 0.2
    avail_h = BODY_BOTTOM - top
    img = min(4.3, avail_h - 0.5)
    rx = MX + img + 0.5
    rw = W_IN - MX - rx
    blocks = [{'label': k, 'text': v} for k, v in d['rows']]
    if d.get('marks'):
        mk = [{'label': None, 'text': m} for m in d['marks']]
        blocks = mk + blocks
    if d.get('caveat'):
        blocks.append({'label': d['caveat'][0], 'text': d['caveat'][1]})
    badge_h = 0.5 if d.get('badge') else 0
    pages = paginate_blocks(blocks, rw, avail_h - badge_h, {'inline': True}, sizes=(26, 24, 22, 20, 18, 16), page_size=18)
    for i, (bl, sz) in enumerate(pages, 1):
        s = ctx.new(spec)
        y = title_block(ctx, s, spec.title, spec.eyebrow, (i, len(pages)))
        add_image(s, spec.image, MX, y, img, img, descr=spec.title)
        add_text(s, MX, y + img + 0.05, img, 0.3, [Para(d['note'], 12, italic=True, color=SUB)])
        ry = y
        if d.get('badge'):
            add_rect(s, rx, ry, rw, 0.42, fill=AMBER_BG, radius=0.1)
            add_text(s, rx + 0.15, ry, rw - 0.3, 0.42, [Para(d['badge'], 16, bold=True, color=AMBER)], anchor='m')
            ry += badge_h
        add_text(s, rx, ry, rw, BODY_BOTTOM - ry, block_paras(bl, sz, {'inline': True}))


def r_pair(ctx, spec):
    d = spec.data
    tip = d['tip']
    title_n = min(2, count_lines(spec.title, 34, CW, True))
    top = 0.36 + 0.32 + title_n * 34 * LH / 72.0 + 0.2
    badge_h = 0.5 if d.get('badge') else 0
    note_h = 0.28
    for tsz in (22, 20, 18):
        tip_h = count_lines(tip, tsz, CW - 0.4) * tsz * LH / 72 + 0.12
        img = BODY_BOTTOM - top - 0.5 - tip_h - badge_h - note_h - 0.1
        if img >= 2.3:
            break
    img = min(img, 3.9)
    s = ctx.new(spec)
    y = title_block(ctx, s, spec.title, spec.eyebrow)
    colw = CW / 2
    for k, key in enumerate(('a', 'b')):
        x = MX + k * colw
        add_image(s, d[key]['image'], x, y, colw - 0.2, img, descr=d[key]['name'])
        add_text(s, x, y + img + 0.02, colw - 0.2, 0.45,
                 [Para(d[key]['name'], 20, bold=True, color=BRAND, align='c')])
    yy = y + img + 0.5
    if d.get('badge'):
        add_rect(s, MX, yy, CW, 0.42, fill=AMBER_BG, radius=0.1)
        add_text(s, MX + 0.15, yy, CW - 0.3, 0.42, [Para(d['badge'], 16, bold=True, color=AMBER)], anchor='m')
        yy += badge_h
    add_rect(s, MX, yy + 0.05, CW, tip_h + 0.05, fill=SOFT, radius=0.1)
    add_text(s, MX + 0.2, yy + 0.1, CW - 0.4, tip_h, [Para(tip, tsz, color=INK)])
    add_text(s, MX, yy + tip_h + 0.14, 6, note_h - 0.04, [Para(d['note'], 12, italic=True, color=SUB)])


def _gallery_pair_fits(items, top):
    for cap_sz in (18, 16):
        colw = CW / len(items)
        cap_h = max(count_lines(i['caption'], cap_sz, colw - 0.3) for i in items) * cap_sz * LH / 72
        cred_h = max(count_lines(i['credit'], 12, colw - 0.3) for i in items) * 12 * LH / 72
        img_h = BODY_BOTTOM - top - cap_h - cred_h - 0.2
        if img_h >= 3.0:
            return cap_sz, cap_h, cred_h, img_h
    return None


def r_gallery(ctx, spec):
    items = spec.data['items']
    title_n = min(2, count_lines(spec.title, 34, CW, True))
    top = 0.36 + 0.32 + title_n * 34 * LH / 72.0 + 0.2
    fit = _gallery_pair_fits(items, top) if len(items) == 2 else None
    if len(items) == 2 and not fit:
        # sarlavhalar uzun: har rasm alohida slaydda (rasm chapda, izoh o'ngda)
        for it in items:
            r_gallery(ctx, dc.Slide('gallery', title=spec.title, eyebrow=spec.eyebrow,
                                    data={'items': [it]}, sources=spec.sources))
        return
    if len(items) == 1:
        it = items[0]
        iw = 7.4
        tx = MX + iw + 0.4
        tw = W_IN - MX - tx
        crh = count_lines(it['credit'], 13, tw) * 13 * LH / 72
        cap = it['caption']
        for csz in (22, 20, 18, 16):
            ch = count_lines(cap, csz, tw) * csz * LH / 72
            if ch + crh + 0.3 <= BODY_BOTTOM - top:
                break
        else:
            # juda uzun izoh: gaplar bo'yicha qisqartiriladi (to'liq matn — izohlar sahifasida)
            csz = 16
            keep = []
            for sent in split_sentences(cap):
                trial = ' '.join(keep + [sent]) + ' …'
                if count_lines(trial, csz, tw) * csz * LH / 72 + crh + 0.3 > BODY_BOTTOM - top:
                    break
                keep.append(sent)
            cap = (' '.join(keep) or cap[:300]) + ' …'
            ch = count_lines(cap, csz, tw) * csz * LH / 72
        s = ctx.new(spec)
        if cap != it['caption']:
            s.notes_slide.notes_text_frame.text = it['caption']
        y = title_block(ctx, s, spec.title, spec.eyebrow)
        add_image(s, it['image'], MX, y, iw, BODY_BOTTOM - y, descr=it['caption'][:250], align='l')
        add_text(s, tx, y, tw, ch + 0.05, [Para(cap, csz, color=INK)])
        add_text(s, tx, y + ch + 0.2, tw, crh + 0.05, [Para(it['credit'], 13, color=SUB)])
        return
    cap_sz, cap_h, cred_h, img_h = fit
    colw = CW / 2
    s = ctx.new(spec)
    y = title_block(ctx, s, spec.title, spec.eyebrow)
    for k, it in enumerate(items):
        x = MX + k * colw
        add_image(s, it['image'], x, y, colw - 0.3, img_h, descr=it['caption'])
        add_text(s, x, y + img_h + 0.08, colw - 0.3, cap_h + 0.05, [Para(it['caption'], cap_sz, color=INK)])
        add_text(s, x, y + img_h + 0.1 + cap_h, colw - 0.3, cred_h + 0.05, [Para(it['credit'], 12, color=SUB)])


LETTERS = 'ABCDEFGH'


def r_qa_q(ctx, spec):
    d = spec.data
    q, opts = d['question'], d['options']
    title_n = 1
    top = 0.36 + 0.32 + title_n * 34 * LH / 72.0 + 0.2
    for qs, os_ in ((30, 24), (28, 22), (26, 20), (24, 18), (22, 18)):
        qh = count_lines(q, qs, CW, True) * qs * LH / 72
        ohs = [max(count_lines(o, os_, CW - 1.3), 1) * os_ * LH / 72 + 0.22 for o in opts]
        if qh + 0.25 + sum(ohs) + 0.1 * len(opts) <= BODY_BOTTOM - top:
            break
    s = ctx.new(spec)
    y = title_block(ctx, s, spec.title, spec.eyebrow)
    add_text(s, MX, y, CW, qh + 0.05, [Para(q, qs, bold=True, color=INK)])
    y += qh + 0.25
    for k, (o, oh) in enumerate(zip(opts, ohs)):
        add_rect(s, MX, y, CW, oh, fill=PAPER, line=LINE, radius=0.1)
        add_rect(s, MX + 0.12, y + (oh - 0.5) / 2, 0.5, 0.5, fill=SOFT, radius=0.25)
        add_text(s, MX + 0.12, y + (oh - 0.5) / 2, 0.5, 0.5,
                 [Para(LETTERS[k], 20, bold=True, color=BRAND, align='c')], anchor='m')
        add_text(s, MX + 0.8, y + 0.02, CW - 1.0, oh - 0.04, [Para(o, os_, color=INK)], anchor='m')
        y += oh + 0.1


def r_qa_a(ctx, spec):
    d = spec.data
    ci = d['correct']
    ans = d['options'][ci]
    blocks = []
    if d.get('explanation'):
        blocks.append({'label': d['labels']['explanation'], 'text': d['explanation']})
    if d.get('basis'):
        blocks.append({'label': d['labels']['basis'], 'text': d['basis']})
    top = 0.36 + 0.32 + 34 * LH / 72.0 + 0.2
    qsz = 18
    qh = count_lines(d['question'], qsz, CW) * qsz * LH / 72
    for asz in (28, 26, 24, 22, 20):
        ah = count_lines(ans, asz, CW - 1.5, True) * asz * LH / 72 + 0.35
        if ah <= 2.6:
            break
    avail = BODY_BOTTOM - top - qh - 0.2 - ah - 0.25
    pages = []   # (bloklar, shrift, ixcham)
    if not blocks:
        pages = [([], 20, False)]
    else:
        pg = paginate_blocks(blocks, CW, avail, {}, sizes=(24, 22, 20, 18), page_size=18, warn=False)
        if len(pg) == 1:
            pages = [(pg[0][0], pg[0][1], False)]
        else:
            # izoh sig'maydi: 1-sahifada savol+javob, keyingilarida ixcham javob + izoh
            avail2 = BODY_BOTTOM - top - 0.9
            pages = [([], 20, False)]
            pages += [(b_, z_, True) for b_, z_ in paginate_blocks(blocks, CW, avail2, {}, sizes=(24, 22, 20),
                                                                  page_size=20)]
    for i, (bl, sz, compact) in enumerate(pages, 1):
        s = ctx.new(spec)
        y = title_block(ctx, s, spec.title, spec.eyebrow, (i, len(pages)))
        if not compact:
            add_text(s, MX, y, CW, qh + 0.05, [Para(d['question'], qsz, italic=True, color=SUB)])
            y += qh + 0.2
            h = ah
            fsz = asz
        else:
            h, fsz = 0.7, 20
        add_rect(s, MX, y, CW, h, fill=BRAND, radius=0.12)
        add_rect(s, MX + 0.2, y + (h - 0.5) / 2, 0.5, 0.5, fill=LIME, radius=0.25)
        add_text(s, MX + 0.2, y + (h - 0.5) / 2, 0.5, 0.5,
                 [Para(LETTERS[ci], 22, bold=True, color=DEEP, align='c')], anchor='m')
        add_text(s, MX + 0.95, y, CW - 1.2, h, [Para(ans if not compact else _clip(ans, 20, CW - 1.2), fsz,
                                                      bold=True, color=WHITE)], anchor='m')
        y += h + 0.25
        if bl:
            add_text(s, MX, y, CW, BODY_BOTTOM - y, block_paras(bl, sz, {}))


def _clip(text, size, width):
    """Bir qatorga sig'adigan qisqartma (ixcham javob satri uchun)."""
    if count_lines(text, size, width, True) <= 1:
        return text
    words = text.split(' ')
    while words and count_lines(' '.join(words) + ' …', size, width, True) > 1:
        words.pop()
    return ' '.join(words) + ' …'


def r_sources(ctx, spec):
    items = [{'label': None, 'text': t} for t in spec.data['items']]
    top = 0.36 + 0.32 + 34 * LH / 72.0 + 0.2
    pages = paginate_blocks(items, CW, BODY_BOTTOM + 0.4 - top, {'numbered': True},
                            sizes=(18, 17, 16, 15, 14), page_size=14)
    for i, (bl, sz) in enumerate(pages, 1):
        s = ctx.new(spec)
        y = title_block(ctx, s, spec.title, '', (i, len(pages)))
        add_text(s, MX, y, CW, BODY_BOTTOM + 0.4 - y, block_paras(bl, sz, {'numbered': True}))
        # raqamlash keyingi sahifada davom etishi uchun boshlang'ich raqam
        if i > 1:
            start = sum(len(p[0]) for p in pages[:i - 1]) + 1
            _renumber(s, start)


def _renumber(slide, start):
    p = _imports()
    qn = p['qn']
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for par in sh.text_frame.paragraphs:
            auto = par._p.find('.//' + qn('a:buAutoNum'))
            if auto is not None:
                auto.set('startAt', str(start))
                break


RENDER = {'title': r_title, 'outline': r_outline, 'text': r_text, 'columns': r_columns,
          'table': r_table, 'card': r_card, 'pair': r_pair, 'gallery': r_gallery,
          'qa_q': r_qa_q, 'qa_a': r_qa_a, 'sources': r_sources}


def render_deck(deck, path):
    ctx = Ctx(deck)
    for spec in deck.slides:
        RENDER[spec.kind](ctx, spec)
    total = len(ctx.pages)
    for n, (s, spec, dark) in enumerate(ctx.pages, 1):
        chrome(ctx, s, spec, n, total, dark)
    os.makedirs(os.path.dirname(path) or '.', exist_ok=True)
    ctx.prs.save(path)
    return total


# ---------------------------------------------------------------- CLI -------

def slug(s):
    t = s.lower()
    for a, b in (('o‘', 'o'), ('g‘', 'g'), ('‘', ''), ('’', ''), ("'", '')):
        t = t.replace(a, b)
    t = re.sub(r'[^a-z0-9]+', '-', t).strip('-')
    return t[:40].strip('-')


def pick_days(repo, topics, all_):
    if all_:
        return list(repo.days.values())
    out = []
    for t in topics:
        m = re.match(r'^d?(\d{1,3})(?:-[PLA])?$', t, re.I)
        if not m or int(m.group(1)) not in repo.days:
            raise SystemExit(f'noma’lum mavzu: {t}')
        d = repo.days[int(m.group(1))]
        if d not in out:
            out.append(d)
    return out


def coverage_line(cov):
    parts = []
    if cov['analytes']:
        parts.append(f'analit {len(cov["analytes"])}')
    if cov['conditions']:
        parts.append(f'holat {len(cov["conditions"])}')
    if cov['reftopics']:
        parts.append(f'jadval {len(cov["reftopics"])}')
    if cov['cells']:
        parts.append(f'hujayra {len(cov["cells"])}')
    if cov['cyto']:
        parts.append(f'sitologiya {len(cov["cyto"])}')
    if cov['atlas']:
        parts.append(f'foto {cov["atlas"]}')
    if cov['oral_reference']:
        parts.append(f'me’yor qatori {cov["oral_reference"]}')
    if cov['books']:
        parts.append(f'kitob {cov["books"]}')
    q = cov['qa']
    if q and any(q.values()):
        parts.append(f'Q&A {q.get("quiz", 0)}+{q.get("test", 0)}+{q.get("oral", 0)}')
    return ', '.join(parts) or '—'


def write_report(path, rows, nodeck, langs, totals, extra=None):
    md = ['# Taqdimotlar qamrovi (mavzu bo‘yicha PPTX)', '',
          'Avtomatik: `python3 tool/slides/make_topic_decks.py --all --dry-run --report docs/DECKS_COVERAGE.md`. '
          'Qo‘lda tahrirlamang.', '',
          '**Ilova ichida:** shu mazmunning hammasi (98 kun mavzusi + amaliy protokollar, uz/ru/en) ilovaning '
          'Ma’ruza bo‘limida qidiruv bilan mavjud: `python3 tool/slides/export_app_decks.py` '
          '`assets/lectures/decks/` va `assets/lectures/index_v2.json` ni qayta yozadi (kontent o‘zgarganda '
          'ishga tushiring; `--check` dolzarbligini tekshiradi). PPTX fayllar repoga qo‘shilmaydi.', '',
          f'Tillar: {", ".join(langs)}. Bitta deck = bitta o‘quv kuni (amaliy + ma’ruza bir xil havolalarga ega, '
          f'`dNNN-P` va `dNNN-L` bitta fayl). Kunlar: {totals["days"]}; deck: {totals["decks"]}; '
          f'kontentsiz (deck yo‘q): {totals["nodeck"]}.', '',
          '## Rasm va jadvallarni mavzuga bog‘lash qoidasi', '',
          '- **Analit/holat kartalari** — `links.analytes` / `links.conditions` (pack.json); har analit: maqsad+fiziologiya, '
          'me’yor/qaror chegaralari (faqat kartada `decision_limits`/`treatment_goals` bo‘lsa), oshishi/pasayishi, preanalitika.',
          '- **Me’yoriy qiymatlar (toifa bankidan)** — mavzuning og‘zaki savollaridagi `reference`/`cutoffs` qatorlari (manbali).',
          '- **Jadvallar** — `links.reference` (anemiya algoritmi, soxta natijalar, sariqlik, jigar sindromlari, najas parazitlari).',
          '- **Leykoformula hujayralari (sxema)** — 7 asosiy hujayra: `/lab/calculators/differential` vositasi yoki '
          'neytrofil/limfotsit/monotsit/eozinofil/bazofil analiti bor mavzularda; reaktiv limfotsit va toksik neytrofil — faqat differential vositali mavzularda; '
          'gipersegmentlangan neytrofil — B12/foliy/MMA analitli yoki sarlavhada “B-12”; blast — sarlavhada “leykoz” (surunkali leykozlardan tashqari); '
          'smudge hujayra — sarlavhada “limfoleykoz”. Adashtiriladigan juftliklar, talqin, surtma texnikasi va misol oraliqlari — differential vositali mavzularda.',
          '- **Sitologiya sxemalari** — `pap-test` → epiteliy hujayralari + Bethesda jadvali; `vaginal-wet-mount` → flora/yallig‘lanish sxemalari; '
          '`pap-test` yoki `hpv-test` → koilotsit, HSIL; surtma olish va bo‘yash — pap-test va vaginal-wet-mount birga bo‘lgan mavzuda. '
          '“Qin surtmasi tozalik darajalari” jadvali manbasi yo‘qligi uchun kiritilmagan.',
          '- **Darslik bo‘yicha qisqa mazmun** — `tool/slides/data/books_extra.json`: ustoz kitoblaridan o‘z so‘zim bilan '
          '(Aripova va b., 2007: d044, d047, d051, d052 — sahifalar bilan; Najmitdinov, 1998: d010, d017, d018 — '
          'elektron matnda sahifa raqami yo‘q, bo‘lim nomi bilan). Yupqa kunlar uchun qo‘shimcha (d001–d006, d009, d010, d019, d020, d032, d033, '
          'd038, d040, d041, d053, d063, d064, d078, d081, d083, d085, d086, d116, d143): Aripov va b., 2006; Yusupova va b., 2022 va Kurbonova, 2023 '
          '(bosma sahifa raqami bilan); Agzamxodjaev va b., 2016; muallifi/yili ko‘rsatilmagan ma’ruza matnlari (shunday belgilangan). '
          'Matn/rasm ko‘chirilmagan; kitobdagi eskirgan yoki o‘zaro zid raqamlar olinmagan. '
          'Kitoblar Drive MCP orqali o‘qildi; Najmitdinov matni dastlab leykoz bobining oxirida (sitokimyo) kesilgan edi, '
          'to‘liq matn keyin .doc orqali tiklandi (d021/d022 uchun kitob mazmuni hali qo‘shilmagan; surunkali mieloleykozning klinik bo‘limi manbada shikastlangan). '
          'Axloq/deontologiya (d001, d004), TORCH (d083, toksoplazmadan tashqari), SOE analizatorlari (d085) va Mindray (d041, d116) '
          'bo‘yicha o‘qiladigan manba topilmadi. DjVu atlaslar, Hodjibekov va Sultonov (o‘qib bo‘lmaydi) olinmadi.',
          '- **Mikroskopik atlas fotolari** — `links.microscopy` (urine / blood / parasites); litsenziya va muallif har slaydda va “Adabiyotlar”da.',
          '- **Savol-javob** — curriculum nomzodlari (yo‘q bo‘lsa ilovadagi kabi analit/toifa mavzusi bo‘yicha): paket mashq savollari '
          '(manbasizlari tashlanadi), toifa testlari (faqat kaliti aniq: verdict `ok`, bitta javob, `unverified` emas), og‘zaki savollar (javob rejasi). '
          f'Chegara: {dc.QA_CAPS}.',
          '- `bloodgroup/card.png` va `spermogram/card.png` — hech bir curriculum mavzusiga bog‘lanmagan (havolalarda yo‘q), shuning uchun ishlatilmadi.', '',
          '## Deck bor mavzular', '',
          '| Kun | Mavzu (uz) | Slayd (PPTX, uz; ~ — taxminiy) | Kontent | Yetishmaydi |', '|---|---|---|---|---|']
    for r in rows:
        md.append(f'| {r["id"]} | {r["title"][:70]} | {r["slides"]} | {r["cov"]} | {"; ".join(r["missing"]) or "—"} |')
    md += ['', '## Deck yo‘q mavzular (kontent topilmadi)', '',
           'Bu kunlarda `links` bo‘sh yoki bog‘langan kontent yo‘q — slayd o‘ylab topilmadi.', '',
           '| Kun | Mavzu (uz) | Sabab |', '|---|---|---|']
    for r in nodeck:
        md.append(f'| {r["id"]} | {r["title"][:70]} | {"; ".join(r["missing"]) or "havolalar bo‘sh"} |')
    md.append('')
    md += extra or []
    with open(path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(md))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--topic', action='append', default=[], help='dNNN yoki dNNN-P/L (takrorlash mumkin)')
    ap.add_argument('--all', action='store_true')
    ap.add_argument('--lang', default='uz', help='uz,ru,en yoki all')
    ap.add_argument('--out', default='decks_out')
    ap.add_argument('--dry-run', action='store_true', help='PPTX yaratmasdan ro‘yxat chiqaradi')
    ap.add_argument('--report', help='qamrov hisoboti (Markdown) yo‘li')
    ap.add_argument('--zip', help='yaratilgan decklarni (birinchi til) zip + README qilish')
    a = ap.parse_args(argv)
    if not a.all and not a.topic:
        ap.error('--topic yoki --all kerak')
    langs = list(dc.LANGS) if a.lang == 'all' else [x.strip() for x in a.lang.split(',')]
    for lg in langs:
        if lg not in dc.LANGS:
            ap.error(f'noma’lum til: {lg}')
    repo = dc.Repo()
    days = pick_days(repo, a.topic, a.all)
    rows, nodeck, made = [], [], []
    counts = {lg: 0 for lg in langs}
    for day in days:
        for lg in langs:
            deck, cov = dc.build_deck(repo, day, lg)
            rec = None
            if lg == langs[0]:
                rec = {'id': day.id, 'title': day.title('uz'), 'cov': coverage_line(cov),
                       'missing': cov['missing'], 'slides': f'~{cov.get("slides", 0)}'}
                (rows if deck else nodeck).append(rec)
            if deck is None:
                if a.dry_run or lg == langs[0]:
                    print(f'{day.id}  —  kontent yo‘q: {"; ".join(cov["missing"]) or "havolalar bo‘sh"}')
                continue
            if a.dry_run:
                if lg == langs[0]:
                    print(f'{day.id}  ~{cov["slides"]:>3} slayd  {coverage_line(cov)}'
                          + (f'  | yetishmaydi: {"; ".join(cov["missing"])}' if cov['missing'] else ''))
                counts[lg] += 1
                continue
            fn = os.path.join(a.out, lg, f'{day.id}_{slug(day.title("uz"))}.pptx')
            n = render_deck(deck, fn)
            if rec is not None:
                rec['slides'] = n
            counts[lg] += 1
            made.append((lg, day, fn, n, deck))
            print(f'{fn}  ({n} slayd)')
    print('decklar:', counts)
    if a.report:
        extra = []
        try:  # amaliy mashg'ulot protokollari (kunga bog'lanmagan) bo'limi
            import practice_decks
            extra = practice_decks.report_section(repo)
        except Exception as e:  # hisobot baribir yoziladi
            print('amaliy protokollar bo‘limi o‘tkazib yuborildi:', e)
        write_report(a.report, rows, nodeck, langs,
                     {'days': len(days), 'decks': len(rows), 'nodeck': len(nodeck)}, extra)
        print('hisobot:', a.report)
    if a.zip and made:
        first = langs[0]
        with zipfile.ZipFile(a.zip, 'w', zipfile.ZIP_DEFLATED) as z:
            lines = ['LabGuide — mavzu taqdimotlari (16:9, televizor uchun)', '',
                     f'Til: {first}. Har fayl — bitta o‘quv kuni (amaliy + ma’ruza birga).', '',
                     'Tuzilishi: sarlavha → dars rejasi → mazmun slaydlari → "Adabiyotlar".',
                     'Har mazmun slaydida pastda qisqa manba qatori (muallif/nashr + bo‘lim yoki sahifa);',
                     'to‘liq ro‘yxat — oxirgi "Adabiyotlar" slaydida. Footer: "LabGuide" va slayd raqami.',
                     'Mikroskopik foto slaydlarida muallif va litsenziya ko‘rsatilgan. Hujayra va sitologiya',
                     'rasmlari — LabGuide sxemalari (mikrofoto emas), slaydda shunday yozilgan.',
                     'Slaydlar tahrirlanadi (PowerPoint/LibreOffice). Shrift: Inter (bo‘lmasa o‘rniga boshqasi tushadi).',
                     'Mazmun ilova kontentidan avtomatik yig‘ilgan; kontenti yo‘q mavzular uchun fayl yo‘q.', '',
                     'Fayl | Mavzu | Slayd | Manbalar | Mazmun', '--- | --- | --- | --- | ---']
            for lg, day, fn, n, deck in made:
                if lg != first:
                    continue
                z.write(fn, os.path.basename(fn))
                lines.append(f'{os.path.basename(fn)} | {day.title(first)} | {n} | {len(deck.sources.entries)} | '
                             f'{coverage_line(deck.coverage)}')
            z.writestr('README.txt', '\n'.join(lines) + '\n')
        print('zip:', a.zip)


if __name__ == '__main__':
    main()
